"""Tabbed Writer: a simple word processor where each tab is a chapter.

The book settings follow Amazon KDP's published guidance:
- eBooks: chapter titles use Heading 1 and are centered, body paragraphs
  have a 0.2 inch first-line indent with no extra space before or after,
  and Word (.docx) or EPUB files are uploaded.
- Paperbacks: standard trim sizes (6 x 9 in is the most common), mirrored
  inside/outside margins, and a print PDF.

Needs: Python 3, PySide6, python-docx (for Word files) and
pyspellchecker (for spell check).
LibreOffice is optional: it's only used to open .odt files.
Works on Linux and Windows.
"""

import base64
import hashlib
import html as html_text
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from PySide6.QtCore import QBuffer, QByteArray, QIODevice, QMarginsF, QRectF, QSettings, QSizeF, QStandardPaths, Qt, QTimer, QUrl
from PySide6.QtGui import (
    QAction,
    QBrush,
    QColor,
    QDesktopServices,
    QFont,
    QIcon,
    QImage,
    QKeySequence,
    QPageLayout,
    QPainter,
    QPageSize,
    QPdfWriter,
    QSyntaxHighlighter,
    QTextBlockFormat,
    QTextCharFormat,
    QTextCursor,
    QTextDocument,
    QTextDocumentFragment,
    QTextDocumentWriter,
    QTextFormat,
    QTextFrameFormat,
    QTextImageFormat,
    QTextLength,
    QTextListFormat,
    QTextTableFormat,
)
from PySide6.QtPrintSupport import QPrintDialog, QPrinter
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QColorDialog,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QDoubleSpinBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QHeaderView,
    QHBoxLayout,
    QInputDialog,
    QListWidget,
    QListWidgetItem,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMenu,
    QMenuBar,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QToolButton,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

APP_VERSION = "1.2"

DOCUMENT_FILE = None
SUGGESTED_NAME = "Untitled.odt"
# True when the open file was made by another program, so saving over it
# could drop things Tabbed Writer doesn't support yet (like images).
FROM_OTHER_PROGRAM = False

# Where your writing is kept. Open and Save as start here.
DOCUMENTS = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DocumentsLocation)
WRITING_FOLDER = Path(DOCUMENTS or Path.home() / "Documents") / "Tabbed Writer"
WRITING_FOLDER.mkdir(parents=True, exist_ok=True)
CONTENT_ROLE = Qt.ItemDataRole.UserRole
WORDS_ROLE = Qt.ItemDataRole.UserRole + 1

# "standard" documents behave like any word processor.
# "ebook" documents follow the KDP / B&N Press formatting rules.
document_kind = "standard"

app_settings = QSettings("TabbedWriter", "TabbedWriter")
DEFAULT_WORDS_PER_PAGE = 250

# A writing folder chosen in Tools → Preferences (for example, inside Dropbox).
DEFAULT_WRITING_FOLDER = WRITING_FOLDER
chosen_folder = str(app_settings.value("writing_folder", "") or "")
if chosen_folder and Path(chosen_folder).is_dir():
    WRITING_FOLDER = Path(chosen_folder)

# Where unsaved work is kept in case the program closes unexpectedly.
RECOVERY_FOLDER = Path(
    QStandardPaths.writableLocation(QStandardPaths.StandardLocation.GenericDataLocation)
    or Path.home()
) / "TabbedWriter" / "recovery"
RECOVERY_FILE = RECOVERY_FOLDER / f"window-{os.getpid()}.twp"

# The same three kinds of font Kindle Create offers (serif, sans serif,
# monospace). Linux uses the free Liberation fonts; Windows and Mac use
# Times New Roman, Arial and Courier New, which they were designed to match.
ON_LINUX = sys.platform.startswith("linux")
LINUX_FONTS = {"Serif": "Liberation Serif", "Sans serif": "Liberation Sans", "Monospace": "Liberation Mono"}
OTHER_FONTS = {"Serif": "Times New Roman", "Sans serif": "Arial", "Monospace": "Courier New"}
FONT_CHOICES = LINUX_FONTS if ON_LINUX else OTHER_FONTS
FONT_LABELS = {family: label for fonts in (LINUX_FONTS, OTHER_FONTS) for label, family in fonts.items()}
DEFAULT_FONT = FONT_CHOICES["Serif"]
DEFAULT_SIZE = 12
SIZE_CHOICES = ["8", "9", "10", "11", "12", "14", "16", "18", "20", "24", "28", "36"]

# KDP: first-line indent of 0.2 inch (Qt measures this at 96 dots per inch).
BOOK_INDENT = 0.2 * 96

# KDP paperback trim sizes (width, height in inches). 6 x 9 is the most common.
TRIM_SIZES = {
    "5 x 8 in": (5.0, 8.0),
    "5.25 x 8 in": (5.25, 8.0),
    "5.5 x 8.5 in": (5.5, 8.5),
    "6 x 9 in": (6.0, 9.0),
    "7 x 10 in": (7.0, 10.0),
    "8.5 x 11 in (Letter)": (8.5, 11.0),
    "8.27 x 11.69 in (A4)": (8.27, 11.69),
}
# Normal documents: Letter paper, 1 inch margins, like most word processors.
STANDARD_PAGE = {"trim": "8.5 x 11 in (Letter)", "top": 1.0, "bottom": 1.0, "inside": 1.0, "outside": 1.0, "page_numbers": False}
# eBooks: KDP's most common paperback size, in case a print copy is wanted later.
DEFAULT_PAGE = {"trim": "6 x 9 in", "top": 0.75, "bottom": 0.75, "inside": 0.75, "outside": 0.5, "page_numbers": False}
page_settings = dict(STANDARD_PAGE)
SCREEN_DPI = 96

# What "Save as" can make.
# Normal documents are kept as .odt (the open standard) or .docx.
# eBooks are kept as .twp, the Tabbed Writer book-project format.
# PDF, EPUB and plain text are always saved as copies.
SAVE_TYPES = {
    "OpenDocument text (*.odt)": ".odt",
    "Word Document (*.docx)": ".docx",
    "Tabbed Writer book (*.twp)": ".twp",
    "EPUB (*.epub)": ".epub",
    "PDF (*.pdf)": ".pdf",
    "Plain text (*.txt)": ".txt",
}
WORKING_FORMATS = (".odt", ".docx", ".twp")


def default_ending():
    return ".twp" if document_kind == "ebook" else ".odt"


# ---------------------------------------------------------------------------
# Tabs and saving
# ---------------------------------------------------------------------------

def make_item(parent, item_data):
    if parent is None:
        item = QTreeWidgetItem(sections, [item_data["title"]])
    else:
        item = QTreeWidgetItem(parent, [item_data["title"]])

    item.setData(0, CONTENT_ROLE, item_data.get("text", ""))
    item.setData(0, WORDS_ROLE, words_in_html(item_data.get("text", "")))

    for child_data in item_data.get("children", []):
        make_item(item, child_data)

    return item


def item_to_data(item):
    return {
        "title": item.text(0),
        "text": item.data(0, CONTENT_ROLE) or "",
        "children": [
            item_to_data(item.child(index))
            for index in range(item.childCount())
        ],
    }


saved_snapshot = ""


def count_words(text):
    return len(text.replace("\ufffc", " ").split())


def words_in_html(html):
    if not html:
        return 0
    holder = QTextDocument()
    holder.setHtml(html)
    return count_words(holder.toPlainText())


def words_per_page():
    try:
        return max(50, int(app_settings.value("words_per_page", DEFAULT_WORDS_PER_PAGE)))
    except (TypeError, ValueError):
        return DEFAULT_WORDS_PER_PAGE


def page_position():
    # Lay the whole document out on pages (like printing it) to find
    # how many pages there are and which page the cursor is on.
    text_width, text_height = page_text_size()
    layout_document = QTextDocument()
    layout_document.setDefaultFont(QFont(DEFAULT_FONT, DEFAULT_SIZE))
    layout_document.setDocumentMargin(0)
    cursor = QTextCursor(layout_document)
    current = sections.currentItem()
    cursor_position = 0
    for index in range(sections.topLevelItemCount()):
        item = sections.topLevelItem(index)
        if index > 0:
            cursor.insertBlock(QTextBlockFormat(), QTextCharFormat())
        if item is current:
            start = cursor.position()
            cursor.insertFragment(QTextDocumentFragment(editor.document()))
            cursor_position = start + editor.textCursor().position()
        else:
            html = item.data(0, CONTENT_ROLE)
            if html:
                cursor.insertFragment(QTextDocumentFragment.fromHtml(html))
    layout_document.setPageSize(QSizeF(text_width, text_height))
    total = max(1, layout_document.pageCount())

    cursor_position = min(cursor_position, layout_document.characterCount() - 1)
    block = layout_document.findBlock(cursor_position)
    y = layout_document.documentLayout().blockBoundingRect(block).y()
    line = block.layout().lineForTextPosition(cursor_position - block.position())
    if line.isValid():
        y += line.y()
    page = min(total, int(y // text_height) + 1)
    return page, total


def words_by_tab():
    current = sections.currentItem()
    counts = []
    for index in range(sections.topLevelItemCount()):
        item = sections.topLevelItem(index)
        if item is current:
            words = count_words(editor.toPlainText())
        else:
            words = item.data(0, WORDS_ROLE) or 0
        counts.append((item.text(0), words))
    return counts


def estimated_pages(total):
    return -(-total // words_per_page()) if total else 0


def plural(number, word):
    return f"{number:,} {word}" + ("" if number == 1 else "s")


def fill_word_table(table, counts):
    table.setRowCount(len(counts))
    for row, (title, words) in enumerate(counts):
        table.setItem(row, 0, QTableWidgetItem(title))
        number = QTableWidgetItem(f"{words:,}")
        number.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        table.setItem(row, 1, number)


def make_word_table():
    table = QTableWidget(0, 2)
    table.setHorizontalHeaderLabels(["Tab", "Words"])
    table.verticalHeader().setVisible(False)
    table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
    table.setSelectionMode(QAbstractItemView.SelectionMode.NoSelection)
    table.setFocusPolicy(Qt.FocusPolicy.NoFocus)
    table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
    table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
    return table


def update_counts():
    counts = words_by_tab()
    total = sum(words for _, words in counts)
    if not word_panel.isHidden():
        fill_word_table(panel_table, counts)
        panel_total.setText(f"Total: {plural(total, 'word')}")
    if document_kind == "ebook":
        # eBooks have no fixed pages, so this is an estimate from words per page.
        count_label.setText(
            f"eBook  ·  {plural(total, 'word')}  ·  about {plural(estimated_pages(total), 'page')}"
        )
    else:
        page, pages = page_position()
        count_label.setText(f"Page {page:,} of {pages:,}  ·  {plural(total, 'word')}")


def show_word_count():
    counts = words_by_tab()
    total = sum(words for _, words in counts)

    dialog = QDialog(window)
    dialog.setWindowTitle("Word count")
    column = QVBoxLayout(dialog)

    table = make_word_table()
    fill_word_table(table, counts)
    column.addWidget(table)

    summary = f"Total: {plural(total, 'word')}"
    if document_kind == "ebook":
        summary += (
            f"\nAbout {plural(estimated_pages(total), 'page')} "
            f"(at {words_per_page():,} words per page)"
        )
    else:
        summary += f"\n{plural(page_position()[1], 'page')}"
    summary_label = QLabel(summary)
    summary_font = summary_label.font()
    summary_font.setBold(True)
    summary_label.setFont(summary_font)
    column.addWidget(summary_label)

    keep_showing = QCheckBox("Show in the tab panel while typing")
    keep_showing.setChecked(show_word_count_action.isChecked())
    keep_showing.toggled.connect(show_word_count_action.setChecked)
    column.addWidget(keep_showing)

    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
    buttons.rejected.connect(dialog.reject)
    column.addWidget(buttons)

    dialog.resize(380, 360)
    dialog.exec()


def update_counts_soon():
    # Wait for a short pause in typing before recounting.
    count_timer.start()


def change_words_per_page():
    number, accepted = QInputDialog.getInt(
        window,
        "eBook words per page",
        "Words per page, for the eBook page estimate.\n"
        "KDP and B&N don't publish a number; 250 is the common rule of thumb.",
        words_per_page(),
        100,
        600,
        10,
    )
    if accepted:
        app_settings.setValue("words_per_page", number)
        update_counts()


DATA_IMAGE = re.compile(r'src="(data:image/[a-z+.-]+;base64,[A-Za-z0-9+/=]+)"')


def image_resource_type():
    return enum_number(QTextDocument.ResourceType.ImageResource)


def register_images(document, html):
    """Make sure a document can show the pictures stored inside its text."""
    for uri in set(DATA_IMAGE.findall(html or "")):
        url = QUrl(uri)
        if document.resource(image_resource_type(), url) is not None:
            continue
        image = QImage.fromData(base64.b64decode(uri.split(",", 1)[1]))
        if not image.isNull():
            document.addResource(image_resource_type(), url, image)


def image_bytes(uri):
    header, data = uri.split(",", 1)
    return header[5:].split(";")[0], base64.b64decode(data)


def editor_html():
    if editor.toPlainText() == "" and editor.document().blockCount() == 1:
        return ""
    return editor.toHtml()


def current_snapshot():
    current = sections.currentItem()
    if current is not None:
        current.setData(0, CONTENT_ROLE, editor_html())
    data = {
        "format": "Tabbed Writer",
        "version": 1,
        "kind": document_kind,
        "page": page_settings,
        "tabs": [
            item_to_data(sections.topLevelItem(index))
            for index in range(sections.topLevelItemCount())
        ],
    }
    return json.dumps(data, indent=2)


def save_document(automatic=False):
    global saved_snapshot, FROM_OTHER_PROGRAM
    if DOCUMENT_FILE is None:
        if not automatic:
            save_as_document()
        return
    if FROM_OTHER_PROGRAM:
        if automatic:
            return
        answer = QMessageBox.question(
            window,
            "Save over this file?",
            f"“{DOCUMENT_FILE.name}” was made in another program. Saving over it may drop "
            "things Tabbed Writer doesn't support yet, like images or comments.\n\n"
            "Save over it anyway? (Choose No to save under a new name instead.)",
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No
            | QMessageBox.StandardButton.Cancel,
        )
        if answer == QMessageBox.StandardButton.No:
            save_as_document()
            return
        if answer != QMessageBox.StandardButton.Yes:
            return
    snapshot = current_snapshot()
    ending = DOCUMENT_FILE.suffix.lower()
    try:
        keep_previous_version(DOCUMENT_FILE, automatic)
    except OSError:
        pass  # a missing backup shouldn't stop the save itself
    if ending == ".twp":
        DOCUMENT_FILE.write_text(snapshot, encoding="utf-8")
    elif not write_copy(str(DOCUMENT_FILE), snapshot=snapshot, quiet=True):
        return
    FROM_OTHER_PROGRAM = False
    saved_snapshot = snapshot
    clear_recovery()
    remember_recent(DOCUMENT_FILE)
    window.statusBar().showMessage("Saved automatically" if automatic else "Saved", 3000)


# ---------------------------------------------------------------------------
# Previous versions
# ---------------------------------------------------------------------------

def versions_to_keep():
    try:
        return max(0, int(app_settings.value("versions_to_keep", 10)))
    except (TypeError, ValueError):
        return 10


def versions_folder_for(path):
    return WRITING_FOLDER / "Previous versions" / path.name


def versions_of(path):
    folder = versions_folder_for(path)
    if not folder.is_dir():
        return []
    return sorted(folder.glob(f"*{path.suffix}"), reverse=True)


def keep_previous_version(path, automatic):
    """Before saving over a file, tuck its last saved state into "Previous versions"."""
    keep = versions_to_keep()
    if keep == 0 or not path.exists():
        return
    existing = versions_of(path)
    if automatic and existing and time.time() - existing[0].stat().st_mtime < 10 * 60:
        return  # automatic saves add at most one version every ten minutes
    folder = versions_folder_for(path)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H-%M-%S")
    target = folder / f"{path.stem} {stamp}{path.suffix}"
    if not target.exists():
        shutil.copy2(path, target)
    for old in versions_of(path)[keep:]:
        old.unlink(missing_ok=True)


def friendly_time(path):
    moment = datetime.fromtimestamp(path.stat().st_mtime)
    hour = moment.strftime("%I").lstrip("0") or "12"
    return moment.strftime(f"%A, %B {moment.day}, %Y at {hour}:%M %p")


def show_version_history():
    if DOCUMENT_FILE is None:
        QMessageBox.information(window, "Version history", "Save this document first. Previous versions are kept each time you save.")
        return
    versions = versions_of(DOCUMENT_FILE)
    if not versions:
        QMessageBox.information(
            window,
            "Version history",
            f"There are no previous versions of “{DOCUMENT_FILE.name}” yet.\n\n"
            "One is kept each time you save over it.",
        )
        return

    dialog = QDialog(window)
    dialog.setWindowTitle(f"Version history: {DOCUMENT_FILE.name}")
    column = QVBoxLayout(dialog)
    column.addWidget(QLabel("Previous versions, newest first:"))
    listing = QListWidget()
    for version in versions:
        entry = QListWidgetItem(friendly_time(version))
        entry.setData(Qt.ItemDataRole.UserRole, str(version))
        listing.addItem(entry)
    listing.setCurrentRow(0)
    column.addWidget(listing)

    def chosen():
        entry = listing.currentItem()
        return Path(entry.data(Qt.ItemDataRole.UserRole)) if entry else None

    def open_copy():
        if chosen():
            start_new_window(str(chosen()))

    def restore():
        version = chosen()
        if version is None:
            return
        answer = QMessageBox.question(
            dialog,
            "Restore this version?",
            "Replace what's in this window with this version?\n\n"
            "Nothing is saved until you save, and the current file is kept as a previous version when you do.",
        )
        if answer == QMessageBox.StandardButton.Yes:
            dialog.accept()
            restore_version(version)

    buttons = QHBoxLayout()
    for label, action in (
        ("Open in new window", open_copy),
        ("Restore this version", restore),
        ("Show folder", lambda: QDesktopServices.openUrl(QUrl.fromLocalFile(str(versions_folder_for(DOCUMENT_FILE))))),
        ("Close", dialog.reject),
    ):
        button = QPushButton(label)
        button.clicked.connect(action)
        buttons.addWidget(button)
    column.addLayout(buttons)
    listing.itemDoubleClicked.connect(lambda item: open_copy())
    dialog.resize(520, 360)
    dialog.exec()


def restore_version(version):
    global DOCUMENT_FILE, FROM_OTHER_PROGRAM, saved_snapshot
    original = DOCUMENT_FILE
    if not load_file(version):
        return
    DOCUMENT_FILE, FROM_OTHER_PROGRAM = original, False
    window.setWindowTitle(f"{original.name} (restored version) - Tabbed Writer")
    saved_snapshot = ""  # counts as unsaved until you save


# ---------------------------------------------------------------------------
# Automatic saving and crash recovery
# ---------------------------------------------------------------------------

def autosave_on():
    return str(app_settings.value("autosave", "true")).lower() == "true"


def autosave_minutes():
    try:
        return max(1, int(app_settings.value("autosave_minutes", 2)))
    except (TypeError, ValueError):
        return 2


def write_recovery():
    data = json.loads(current_snapshot())
    data["recovery_for"] = str(DOCUMENT_FILE) if DOCUMENT_FILE else ""
    data["from_other_program"] = FROM_OTHER_PROGRAM
    data["saved_at"] = datetime.now().isoformat(timespec="seconds")
    RECOVERY_FOLDER.mkdir(parents=True, exist_ok=True)
    temporary = RECOVERY_FILE.with_suffix(".tmp")
    temporary.write_text(json.dumps(data), encoding="utf-8")
    temporary.replace(RECOVERY_FILE)


def clear_recovery():
    try:
        RECOVERY_FILE.unlink(missing_ok=True)
    except OSError:
        pass


minutes_since_autosave = 0


def autosave_tick():
    """Runs every minute."""
    global minutes_since_autosave
    minutes_since_autosave += 1
    if QApplication.activeModalWidget() is not None:
        return  # don't save in the middle of a question or a Save window
    if current_snapshot() == saved_snapshot:
        clear_recovery()
        return
    if (
        autosave_on()
        and DOCUMENT_FILE is not None
        and not FROM_OTHER_PROGRAM
        and minutes_since_autosave >= autosave_minutes()
    ):
        minutes_since_autosave = 0
        save_document(automatic=True)
        if current_snapshot() == saved_snapshot:
            return
    try:
        write_recovery()
    except OSError:
        pass


def process_running(pid):
    if pid == os.getpid():
        return True
    if sys.platform.startswith("win"):
        import ctypes
        kernel = ctypes.windll.kernel32
        handle = kernel.OpenProcess(0x1000, False, pid)  # ask about it, nothing more
        if not handle:
            return False
        code = ctypes.c_ulong()
        found = kernel.GetExitCodeProcess(handle, ctypes.byref(code))
        kernel.CloseHandle(handle)
        return bool(found) and code.value == 259  # still running
    try:
        os.kill(pid, 0)  # on Linux this only checks; it sends nothing
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True


def leftover_recoveries():
    if not RECOVERY_FOLDER.is_dir():
        return []
    found = []
    for path in RECOVERY_FOLDER.glob("window-*.twp"):
        try:
            pid = int(path.stem.split("-", 1)[1])
        except (IndexError, ValueError):
            continue
        if not process_running(pid):
            found.append(path)
    return found


def check_for_recovery():
    for path in leftover_recoveries():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            path.unlink(missing_ok=True)
            continue
        name = Path(data.get("recovery_for") or "Untitled").name
        when = data.get("saved_at", "").replace("T", " at ")
        box = QMessageBox(window)
        box.setWindowTitle("Recover unsaved work?")
        box.setText(
            f"Tabbed Writer closed before “{name}” was saved.\n"
            f"Unsaved work from {when} can be recovered."
        )
        box.addButton("Recover", QMessageBox.ButtonRole.AcceptRole)
        box.addButton("Discard", QMessageBox.ButtonRole.DestructiveRole)
        box.addButton("Ask me later", QMessageBox.ButtonRole.RejectRole)
        box.exec()
        clicked = box.clickedButton()
        role = box.buttonRole(clicked) if clicked is not None else QMessageBox.ButtonRole.RejectRole
        if role == QMessageBox.ButtonRole.DestructiveRole:
            path.unlink(missing_ok=True)
        elif role == QMessageBox.ButtonRole.AcceptRole:
            pristine = DOCUMENT_FILE is None and current_snapshot() == saved_snapshot
            if pristine:
                open_document(str(path))
            else:
                start_new_window(str(path))
                return  # the new window takes care of the rest


def recover_from(path):
    """Load a recovery file into this window, pointing back at the original document."""
    global DOCUMENT_FILE, FROM_OTHER_PROGRAM, SUGGESTED_NAME
    data = json.loads(path.read_text(encoding="utf-8"))
    apply_saved_data(data)
    original = data.get("recovery_for") or ""
    DOCUMENT_FILE = Path(original) if original else None
    FROM_OTHER_PROGRAM = bool(data.get("from_other_program"))
    if DOCUMENT_FILE is None:
        SUGGESTED_NAME = "Recovered" + default_ending()
    path.unlink(missing_ok=True)
    return DOCUMENT_FILE.name if DOCUMENT_FILE else "Untitled"


# ---------------------------------------------------------------------------
# Recent files
# ---------------------------------------------------------------------------

def recent_files():
    files = app_settings.value("recent_files", [])
    if isinstance(files, str):
        files = [files]
    return [name for name in (files or []) if name]


def remember_recent(path):
    if path is None:
        return
    name = str(Path(path).resolve())
    files = [name] + [other for other in recent_files() if other != name]
    app_settings.setValue("recent_files", files[:10])


def fill_recent_menu():
    recent_menu.clear()
    files = [name for name in recent_files() if Path(name).exists()]
    if not files:
        empty = recent_menu.addAction("(No recent files)")
        empty.setEnabled(False)
        return
    for name in files:
        path = Path(name)
        action = recent_menu.addAction(f"{path.name}    —  {path.parent}")
        action.triggered.connect(lambda checked=False, n=name: open_document(n))
    recent_menu.addSeparator()
    recent_menu.addAction("Clear list", lambda: app_settings.setValue("recent_files", []))


# ---------------------------------------------------------------------------
# Preferences
# ---------------------------------------------------------------------------

def show_preferences():
    global WRITING_FOLDER
    dialog = QDialog(window)
    dialog.setWindowTitle("Preferences")
    form = QFormLayout(dialog)

    folder_label = QLabel(str(WRITING_FOLDER))
    folder_label.setWordWrap(True)
    new_folder = {"path": WRITING_FOLDER}

    def pick_folder():
        picked = QFileDialog.getExistingDirectory(dialog, "Choose your writing folder", str(new_folder["path"]))
        if picked:
            new_folder["path"] = Path(picked)
            folder_label.setText(picked)

    def default_folder():
        new_folder["path"] = DEFAULT_WRITING_FOLDER
        folder_label.setText(str(DEFAULT_WRITING_FOLDER))

    folder_row = QHBoxLayout()
    folder_row.addWidget(folder_label, 1)
    change = QPushButton("Change…")
    change.clicked.connect(pick_folder)
    reset = QPushButton("Default")
    reset.clicked.connect(default_folder)
    folder_row.addWidget(change)
    folder_row.addWidget(reset)
    form.addRow("Writing folder:", folder_row)
    form.addRow(QLabel(
        "Open and Save as start here, and previous versions are kept here.\n"
        "Choosing a folder inside Dropbox keeps everything backed up."
    ))

    autosave_box = QCheckBox("Save automatically every")
    autosave_box.setChecked(autosave_on())
    minutes_box = QSpinBox()
    minutes_box.setRange(1, 60)
    minutes_box.setValue(autosave_minutes())
    minutes_box.setSuffix(" minutes")
    autosave_row = QHBoxLayout()
    autosave_row.addWidget(autosave_box)
    autosave_row.addWidget(minutes_box)
    autosave_row.addStretch()
    form.addRow("Autosave:", autosave_row)
    form.addRow(QLabel(
        "Only documents that already have a name are saved automatically.\n"
        "Unsaved work is always kept for recovery if the program closes unexpectedly."
    ))

    versions_box = QSpinBox()
    versions_box.setRange(0, 100)
    versions_box.setValue(versions_to_keep())
    versions_box.setSuffix(" per document")
    versions_box.setSpecialValueText("None")
    form.addRow("Keep previous versions:", versions_box)

    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
    buttons.accepted.connect(dialog.accept)
    buttons.rejected.connect(dialog.reject)
    form.addRow(buttons)

    if not dialog.exec():
        return
    WRITING_FOLDER = new_folder["path"]
    WRITING_FOLDER.mkdir(parents=True, exist_ok=True)
    app_settings.setValue(
        "writing_folder", "" if WRITING_FOLDER == DEFAULT_WRITING_FOLDER else str(WRITING_FOLDER)
    )
    app_settings.setValue("autosave", autosave_box.isChecked())
    app_settings.setValue("autosave_minutes", minutes_box.value())
    app_settings.setValue("versions_to_keep", versions_box.value())


def ok_to_discard():
    if current_snapshot() == saved_snapshot:
        return True
    answer = QMessageBox.question(
        window,
        "Unsaved changes",
        "Save changes to this document first?",
        QMessageBox.StandardButton.Save
        | QMessageBox.StandardButton.Discard
        | QMessageBox.StandardButton.Cancel,
    )
    if answer == QMessageBox.StandardButton.Save:
        save_document()
        return current_snapshot() == saved_snapshot
    return answer == QMessageBox.StandardButton.Discard


def page_text_size():
    width, height = TRIM_SIZES.get(page_settings["trim"], TRIM_SIZES[STANDARD_PAGE["trim"]])
    text_width = (width - page_settings["inside"] - page_settings["outside"]) * SCREEN_DPI
    text_height = (height - page_settings["top"] - page_settings["bottom"]) * SCREEN_DPI
    return max(text_width, 100), max(text_height, 100)


def apply_page_look():
    # The writing area looks like a sheet of paper, using the Page setup size and margins.
    width, _ = TRIM_SIZES.get(page_settings["trim"], TRIM_SIZES[STANDARD_PAGE["trim"]])
    root = editor.document().rootFrame()
    frame_format = root.frameFormat()
    frame_format.setLeftMargin(page_settings["inside"] * SCREEN_DPI)
    frame_format.setRightMargin(page_settings["outside"] * SCREEN_DPI)
    frame_format.setTopMargin(page_settings["top"] * SCREEN_DPI)
    frame_format.setBottomMargin(page_settings["bottom"] * SCREEN_DPI)
    root.setFrameFormat(frame_format)
    editor.setMaximumWidth(int(width * SCREEN_DPI) + editor.verticalScrollBar().sizeHint().width())


def use_book_paragraph_if_empty():
    # In eBooks, new empty tabs start with the KDP paragraph style (0.2 inch indent).
    if document_kind != "ebook":
        return
    if editor.toPlainText() == "" and editor.document().blockCount() == 1:
        book_paragraph = QTextBlockFormat()
        book_paragraph.setTextIndent(BOOK_INDENT)
        book_paragraph.setTopMargin(0)
        book_paragraph.setBottomMargin(0)
        editor.textCursor().setBlockFormat(book_paragraph)


def switch_section(current, previous):
    if previous is not None:
        previous.setData(0, CONTENT_ROLE, editor_html())
        previous.setData(0, WORDS_ROLE, count_words(editor.toPlainText()))

    if current is not None:
        register_images(editor.document(), current.data(0, CONTENT_ROLE))
        editor.setHtml(current.data(0, CONTENT_ROLE) or "")
        apply_page_look()
        use_book_paragraph_if_empty()
    update_counts()


def add_main_tab():
    if document_kind == "ebook":
        label, suggested = "Chapter name:", "Chapter"
    else:
        label, suggested = "Tab name:", "Tab"
    title, accepted = QInputDialog.getText(
        window, "Add tab", label, text=f"{suggested} {sections.topLevelItemCount() + 1}"
    )
    if not accepted or not title.strip():
        return
    item = QTreeWidgetItem(sections, [title.strip()])
    item.setData(0, CONTENT_ROLE, "")
    item.setData(0, WORDS_ROLE, 0)
    sections.setCurrentItem(item)


def rename_section():
    item = sections.currentItem()
    if item is None:
        return

    title, accepted = QInputDialog.getText(
        window,
        "Rename tab",
        "Tab name:",
        text=item.text(0),
    )
    if accepted and title.strip():
        item.setText(0, title.strip())
        update_counts_soon()


def remove_section():
    item = sections.currentItem()
    if item is None:
        return

    if sections.topLevelItemCount() <= 1:
        QMessageBox.information(window, "Delete tab", "A document needs at least one tab.")
        return

    answer = QMessageBox.question(
        window,
        "Delete tab",
        f"Delete “{item.text(0)}” and everything written in it?",
        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
    )
    if answer != QMessageBox.StandardButton.Yes:
        return

    index = sections.indexOfTopLevelItem(item)

    sections.blockSignals(True)
    sections.takeTopLevelItem(index)
    count = sections.topLevelItemCount()
    next_item = sections.topLevelItem(min(index, count - 1))
    sections.setCurrentItem(next_item)
    sections.blockSignals(False)

    register_images(editor.document(), next_item.data(0, CONTENT_ROLE))
    editor.setHtml(next_item.data(0, CONTENT_ROLE) or "")
    apply_page_look()
    use_book_paragraph_if_empty()
    update_counts()


def fill_tabs(tabs):
    sections.blockSignals(True)
    sections.clear()
    for tab in tabs:
        make_item(None, tab)
    if sections.topLevelItemCount() == 0:
        make_item(None, {"title": "Tab 1", "text": ""})
    sections.blockSignals(False)
    editor.clear()
    sections.setCurrentItem(sections.topLevelItem(0))


def read_page_settings(saved_page):
    settings = dict(STANDARD_PAGE)
    for key in STANDARD_PAGE:
        if key in saved_page:
            settings[key] = saved_page[key]
    if settings["trim"] not in TRIM_SIZES:
        settings["trim"] = STANDARD_PAGE["trim"]
    return settings


def set_document_kind(kind):
    global document_kind
    document_kind = kind
    ebook_action.blockSignals(True)
    ebook_action.setChecked(kind == "ebook")
    ebook_action.blockSignals(False)


def load_document(path):
    apply_saved_data(json.loads(path.read_text(encoding="utf-8")))


def apply_saved_data(saved):
    global page_settings
    if isinstance(saved, list):
        tabs = saved
        page_settings = dict(STANDARD_PAGE)
        set_document_kind("standard")
    else:
        tabs = saved.get("tabs", [])
        page_settings = read_page_settings(saved.get("page", {}))
        set_document_kind("ebook" if saved.get("kind") == "ebook" else "standard")
    apply_page_look()
    fill_tabs(tabs)


# ---------------------------------------------------------------------------
# LibreOffice (optional: only used to open .odt files)
# ---------------------------------------------------------------------------

def find_libreoffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for place in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if Path(place).exists():
            return place
    return None


def libreoffice_convert(source, target_format, out_dir):
    office = find_libreoffice()
    if office is None:
        raise RuntimeError(
            "Opening .odt files uses LibreOffice, which isn't installed on this computer.\n\n"
            "You can get it free from libreoffice.org, or save the file as a Word "
            "document (.docx) in the program that made it, then open that."
        )
    QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
    try:
        with tempfile.TemporaryDirectory() as profile:
            subprocess.run(
                [
                    office,
                    "--headless",
                    f"-env:UserInstallation={Path(profile).as_uri()}",
                    "--convert-to",
                    target_format,
                    "--outdir",
                    str(out_dir),
                    str(source),
                ],
                capture_output=True,
                timeout=180,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
    finally:
        QApplication.restoreOverrideCursor()


# ---------------------------------------------------------------------------
# Book settings inside .odt files
# ---------------------------------------------------------------------------

ODF_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "dc": "http://purl.org/dc/elements/1.1/",
    "meta": "urn:oasis:names:tc:opendocument:xmlns:meta:1.0",
    "manifest": "urn:oasis:names:tc:opendocument:xmlns:manifest:1.0",
}


def odf_tag(prefix, name):
    return f"{{{ODF_NS[prefix]}}}{name}"


def read_xml(data):
    for prefix, uri in ODF_NS.items():
        ET.register_namespace(prefix, uri)
    for _, (prefix, uri) in ET.iterparse(io.BytesIO(data), events=["start-ns"]):
        if prefix:
            ET.register_namespace(prefix, uri)
    return ET.fromstring(data)


def xml_bytes(root):
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root)


def odf_child(parent, prefix, name, before=None):
    found = parent.find(odf_tag(prefix, name))
    if found is None:
        found = ET.Element(odf_tag(prefix, name))
        later = [parent.find(odf_tag(p, n)) for p, n in (before or [])]
        later = [element for element in later if element is not None]
        if later:
            parent.insert(list(parent).index(later[0]), found)
        else:
            parent.append(found)
    return found


def fix_odt_styles(root, page):
    styles = odf_child(root, "office", "styles",
                       before=[("office", "automatic-styles"), ("office", "master-styles")])
    automatic = odf_child(root, "office", "automatic-styles", before=[("office", "master-styles")])
    masters = odf_child(root, "office", "master-styles")

    names = {element.get(odf_tag("style", "name")) for element in styles}
    if "Heading_20_1" not in names:
        heading = ET.SubElement(styles, odf_tag("style", "style"), {
            odf_tag("style", "name"): "Heading_20_1",
            odf_tag("style", "display-name"): "Heading 1",
            odf_tag("style", "family"): "paragraph",
            odf_tag("style", "default-outline-level"): "1",
            odf_tag("style", "class"): "text",
        })
        ET.SubElement(heading, odf_tag("style", "paragraph-properties"), {
            odf_tag("fo", "text-align"): "center",
            odf_tag("fo", "margin-top"): "0in",
            odf_tag("fo", "margin-bottom"): "0.25in",
            odf_tag("fo", "keep-with-next"): "always",
        })
        ET.SubElement(heading, odf_tag("style", "text-properties"), {
            odf_tag("fo", "font-size"): "18pt",
            odf_tag("fo", "font-weight"): "bold",
        })

    for old in automatic.findall(odf_tag("style", "page-layout")):
        if old.get(odf_tag("style", "name")) == "TWPageLayout":
            automatic.remove(old)
    width, height = TRIM_SIZES.get(page["trim"], TRIM_SIZES[DEFAULT_PAGE["trim"]])
    layout = ET.SubElement(automatic, odf_tag("style", "page-layout"), {
        odf_tag("style", "name"): "TWPageLayout",
        odf_tag("style", "page-usage"): "mirrored",
    })
    ET.SubElement(layout, odf_tag("style", "page-layout-properties"), {
        odf_tag("fo", "page-width"): f"{width}in",
        odf_tag("fo", "page-height"): f"{height}in",
        odf_tag("style", "print-orientation"): "portrait",
        odf_tag("fo", "margin-top"): f"{page['top']}in",
        odf_tag("fo", "margin-bottom"): f"{page['bottom']}in",
        odf_tag("fo", "margin-left"): f"{page['inside']}in",
        odf_tag("fo", "margin-right"): f"{page['outside']}in",
    })

    master_pages = masters.findall(odf_tag("style", "master-page"))
    if not master_pages:
        master_pages = [ET.SubElement(masters, odf_tag("style", "master-page"),
                                      {odf_tag("style", "name"): "Standard"})]
    for master in master_pages:
        master.set(odf_tag("style", "page-layout-name"), "TWPageLayout")


def fix_odt_headings(root):
    automatic = root.find(odf_tag("office", "automatic-styles"))
    by_name = {}
    if automatic is not None:
        by_name = {element.get(odf_tag("style", "name")): element for element in automatic}
    for heading in root.iter(odf_tag("text", "h")):
        if heading.get(odf_tag("text", "outline-level")) != "1":
            continue
        style = by_name.get(heading.get(odf_tag("text", "style-name")))
        if style is None:
            heading.set(odf_tag("text", "style-name"), "Heading_20_1")
        elif style.get(odf_tag("style", "parent-style-name")) is None:
            style.set(odf_tag("style", "parent-style-name"), "Heading_20_1")


TAB_DATA_PATH = "TabbedWriter/tabs.json"


def tuck_tab_data(file_name, snapshot):
    """Store the tabs inside an .odt file, so Tabbed Writer can reopen it exactly.
    Other programs ignore this part. A fingerprint of the visible text is kept too,
    so if the file is edited elsewhere, the stored tabs aren't used."""
    with zipfile.ZipFile(file_name) as source:
        parts = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    contents = dict(parts)
    data = json.loads(snapshot)
    data["content_sha256"] = hashlib.sha256(contents["content.xml"]).hexdigest()
    manifest = contents["META-INF/manifest.xml"]
    if TAB_DATA_PATH.encode() not in manifest:
        manifest = manifest.replace(
            b"</manifest:manifest>",
            f' <manifest:file-entry manifest:media-type="application/json" '
            f'manifest:full-path="{TAB_DATA_PATH}"/>\n</manifest:manifest>'.encode(),
        )
    with zipfile.ZipFile(file_name, "w") as target:
        for name, content in parts:
            if name == TAB_DATA_PATH:
                continue
            if name == "META-INF/manifest.xml":
                content = manifest
            kind = zipfile.ZIP_STORED if name == "mimetype" else zipfile.ZIP_DEFLATED
            target.writestr(name, content, compress_type=kind)
        target.writestr(TAB_DATA_PATH, json.dumps(data, indent=2), compress_type=zipfile.ZIP_DEFLATED)


def read_tab_data(file_name):
    """Return the tabs stored inside an .odt by Tabbed Writer, or None."""
    try:
        with zipfile.ZipFile(file_name) as source:
            if TAB_DATA_PATH not in source.namelist():
                return None
            data = json.loads(source.read(TAB_DATA_PATH))
            fingerprint = hashlib.sha256(source.read("content.xml")).hexdigest()
    except (zipfile.BadZipFile, KeyError, ValueError):
        return None
    if data.get("content_sha256") != fingerprint:
        return None
    return data


def finish_odt(file_name, page, title):
    """Add the page setup, a real Heading 1 style, and the book title to an .odt file."""
    with zipfile.ZipFile(file_name) as source:
        parts = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    contents = dict(parts)
    added = []

    if "styles.xml" in contents:
        styles_root = read_xml(contents["styles.xml"])
    else:
        read_xml(b"<x/>")
        styles_root = ET.Element(odf_tag("office", "document-styles"), {odf_tag("office", "version"): "1.2"})
        added.append("styles.xml")
    fix_odt_styles(styles_root, page)
    contents["styles.xml"] = xml_bytes(styles_root)

    content_root = read_xml(contents["content.xml"])
    fix_odt_headings(content_root)
    contents["content.xml"] = xml_bytes(content_root)

    if "meta.xml" not in contents:
        meta_root = ET.Element(odf_tag("office", "document-meta"), {odf_tag("office", "version"): "1.2"})
        meta = ET.SubElement(meta_root, odf_tag("office", "meta"))
        ET.SubElement(meta, odf_tag("dc", "title")).text = title
        contents["meta.xml"] = xml_bytes(meta_root)
        added.append("meta.xml")

    if added:
        entries = "".join(
            f' <manifest:file-entry manifest:media-type="text/xml" manifest:full-path="{name}"/>\n'
            for name in added
        ).encode()
        contents["META-INF/manifest.xml"] = contents["META-INF/manifest.xml"].replace(
            b"</manifest:manifest>", entries + b"</manifest:manifest>"
        )

    order = [name for name, _ in parts] + added
    with zipfile.ZipFile(file_name, "w") as target:
        for name in order:
            if name == "mimetype":
                target.writestr(name, contents[name], compress_type=zipfile.ZIP_STORED)
            else:
                target.writestr(name, contents[name], compress_type=zipfile.ZIP_DEFLATED)


# ---------------------------------------------------------------------------
# Word and EPUB files, written directly (no LibreOffice needed)
#
# A "book" here is a plain list of items:
#   {"type": "paragraph", "heading": 0 or 1, "align": "left"/"center"/"right"/"justify",
#    "indent": True/False, "page_break": True/False,
#    "runs": [{"text", "bold", "italic", "underline", "color", "highlight", "font", "size"}]}
#   {"type": "table", "border": "#rrggbb" or None,
#    "rows": [[{"background": "#rrggbb" or None, "paragraphs": [paragraph, ...]}, ...], ...]}
# ---------------------------------------------------------------------------

WORD_FONTS = {
    "Liberation Serif": "Times New Roman",
    "Liberation Sans": "Arial",
    "Liberation Mono": "Courier New",
}


GENERIC_FONTS = {
    "Liberation Serif": "serif",
    "Times New Roman": "serif",
    "Liberation Sans": "sans-serif",
    "Arial": "sans-serif",
    "Liberation Mono": "monospace",
    "Courier New": "monospace",
}


def word_font(family):
    return WORD_FONTS.get(family, family)


# ----- Word (.docx) --------------------------------------------------------

def write_docx(book, file_name, page, trim_sizes, title, ebook):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    HYPERLINK = RELATIONSHIP_TYPE.HYPERLINK

    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    document = Document()
    document.core_properties.title = title
    document.core_properties.category = "Tabbed Writer eBook" if ebook else "Tabbed Writer"

    width, height = trim_sizes[page["trim"]]
    section = document.sections[0]
    section.page_width = Inches(width)
    section.page_height = Inches(height)
    section.top_margin = Inches(page["top"])
    section.bottom_margin = Inches(page["bottom"])
    section.left_margin = Inches(page["inside"])
    section.right_margin = Inches(page["outside"])
    if ebook:
        mirror = OxmlElement("w:mirrorMargins")
        document.settings.element.append(mirror)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0 if ebook else 6)

    heading_style = document.styles["Heading 1"]
    heading_style.font.size = Pt(18)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.font.name = "Times New Roman"
    heading_style.paragraph_format.space_before = Pt(0)
    heading_style.paragraph_format.space_after = Pt(18)
    for level, size in ((2, 16), (3, 14)):
        smaller = document.styles[f"Heading {level}"]
        smaller.font.size = Pt(size)
        smaller.font.bold = True
        smaller.font.color.rgb = RGBColor(0, 0, 0)
        smaller.font.name = "Times New Roman"
        smaller.paragraph_format.space_before = Pt(12)
        smaller.paragraph_format.space_after = Pt(6)

    def shade(properties, color):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), color.lstrip("#").upper())
        properties.append(shading)

    def list_style(kind, level):
        base = "List Bullet" if kind == "bullet" else "List Number"
        name = base if level <= 1 else f"{base} {min(level, 3)}"
        try:
            return document.styles[name]
        except KeyError:
            return document.styles[base]

    def fill_paragraph(paragraph, item):
        level = item.get("heading", 0)
        if level:
            paragraph.style = document.styles[f"Heading {min(level, 3)}"]
        elif item.get("list"):
            paragraph.style = list_style(item["list"]["kind"], item["list"]["level"])
        paragraph.alignment = alignments.get(item["align"], WD_ALIGN_PARAGRAPH.LEFT)
        if item["indent"] and not item.get("list"):
            paragraph.paragraph_format.first_line_indent = Inches(0.2)
        if item.get("left") and not item.get("list"):
            paragraph.paragraph_format.left_indent = Inches(0.5 * item["left"])
        if item.get("line") and item["line"] != 100:
            paragraph.paragraph_format.line_spacing = item["line"] / 100
        if item["page_break"]:
            paragraph.paragraph_format.page_break_before = True
        for run_data in item["runs"]:
            if run_data.get("image"):
                try:
                    _, picture = image_bytes(run_data["image"])
                    width = run_data.get("width") or 0
                    paragraph.add_run().add_picture(
                        io.BytesIO(picture), width=Inches(width / 96) if width else None
                    )
                except Exception:
                    pass  # a picture Word can't read is left out rather than failing the save
                continue
            pieces = run_data["text"].split(" ")
            for number, piece in enumerate(pieces):
                if number:
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                if not piece:
                    continue
                run = paragraph.add_run(piece)
                run.bold = run_data["bold"] or None
                run.italic = run_data["italic"] or None
                run.underline = run_data["underline"] or None
                if run_data["font"]:
                    run.font.name = word_font(run_data["font"])
                if run_data["size"]:
                    run.font.size = Pt(run_data["size"])
                if run_data["color"]:
                    run.font.color.rgb = RGBColor.from_string(run_data["color"].lstrip("#").upper())
                if run_data["highlight"]:
                    shade(run._element.get_or_add_rPr(), run_data["highlight"])
                if run_data.get("link"):
                    link_id = paragraph.part.relate_to(run_data["link"], HYPERLINK, is_external=True)
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("r:id"), link_id)
                    run._r.addprevious(link)
                    link.append(run._r)
                    if not run_data["color"]:
                        run.font.color.rgb = RGBColor(0x11, 0x55, 0xCC)
                    run.underline = True
                if run_data.get("strike"):
                    run.font.strike = True
                if run_data.get("vertical") == "super":
                    run.font.superscript = True
                elif run_data.get("vertical") == "sub":
                    run.font.subscript = True

    def set_table_borders(table, color):
        table_properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            line = OxmlElement(f"w:{side}")
            line.set(qn("w:val"), "single")
            line.set(qn("w:sz"), "4")
            line.set(qn("w:space"), "0")
            line.set(qn("w:color"), (color or "#000000").lstrip("#").upper())
            borders.append(line)
        table_properties.append(borders)

    if page.get("page_numbers"):
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        number_run = OxmlElement("w:r")
        number_text = OxmlElement("w:t")
        number_text.text = "1"
        number_run.append(number_text)
        field.append(number_run)
        footer._p.append(field)

    for item in book:
        if item["type"] == "rule":
            ruled = document.add_paragraph()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(key), value)
            borders.append(bottom)
            ruled._p.get_or_add_pPr().append(borders)
            continue
        if item["type"] == "paragraph":
            fill_paragraph(document.add_paragraph(), item)
            continue
        rows = item["rows"]
        columns = max((len(row) for row in rows), default=0)
        if not columns:
            continue
        table = document.add_table(rows=len(rows), cols=columns)
        set_table_borders(table, item["border"])
        for row_number, row in enumerate(rows):
            for column_number, cell_data in enumerate(row):
                cell = table.cell(row_number, column_number)
                if cell_data["background"]:
                    shade(cell._tc.get_or_add_tcPr(), cell_data["background"])
                paragraphs = cell_data["paragraphs"] or []
                for number, paragraph_data in enumerate(paragraphs):
                    paragraph = cell.paragraphs[0] if number == 0 else cell.add_paragraph()
                    fill_paragraph(paragraph, dict(paragraph_data, indent=False, page_break=False))
        document.add_paragraph()

    # Word documents always start with one empty paragraph; remove it.
    first = document.paragraphs[0]
    if not first.text and len(document.paragraphs) > 1:
        first._element.getparent().remove(first._element)

    document.save(file_name)


# ----- EPUB ----------------------------------------------------------------

EPUB_CSS = """body { font-family: serif; margin: 0 5%; }
p { margin: 0 0 0.6em 0; }
p.book { margin: 0; text-indent: 1.2em; }
p.noindent { text-indent: 0; }
h1 { font-size: 1.5em; font-weight: bold; margin: 1em 0 1em 0; }
h2 { font-size: 1.25em; font-weight: bold; margin: 1em 0 0.5em 0; }
h3 { font-size: 1.1em; font-weight: bold; margin: 1em 0 0.5em 0; }
h1.center, h2.center, h3.center { text-align: center; }
p.li { text-indent: 0; margin-top: 0; margin-bottom: 0.2em; }
hr { border: none; border-top: 1px solid #999; margin: 1em 0; }
.center { text-align: center; }
.right { text-align: right; }
.justify { text-align: justify; }
.u { text-decoration: underline; }
table { border-collapse: collapse; width: 100%; margin: 0.6em 0; }
td { border: 1px solid #000; padding: 0.25em; vertical-align: top; }
td p { margin: 0; text-indent: 0; }
"""


def escape(text):
    return html_text.escape(text, quote=True)


def epub_runs(runs):
    out = []
    for run in runs:
        if run.get("image_file"):
            width = run.get("width") or 0
            size = f' style="max-width: 100%; width: {width:g}px"' if width else ' style="max-width: 100%"'
            out.append(f'<img src="{run["image_file"]}" alt=""{size}/>')
            continue
        if run.get("image"):
            continue
        text = "<br/>".join(escape(piece) for piece in run["text"].split(" "))
        if not text:
            continue
        styles = []
        if run["color"]:
            styles.append(f"color: {run['color']}")
        if run["highlight"]:
            styles.append(f"background-color: {run['highlight']}")
        generic = GENERIC_FONTS.get(run["font"] or "", "serif")
        if generic != "serif":
            styles.append(f"font-family: {generic}")
        if run["underline"]:
            text = f'<span class="u">{text}</span>'
        if run.get("strike"):
            text = f"<s>{text}</s>"
        if run.get("vertical") == "super":
            text = f"<sup>{text}</sup>"
        elif run.get("vertical") == "sub":
            text = f"<sub>{text}</sub>"
        if run["italic"]:
            text = f"<em>{text}</em>"
        if run["bold"]:
            text = f"<strong>{text}</strong>"
        if styles:
            text = f'<span style="{"; ".join(styles)}">{text}</span>'
        if run.get("link"):
            text = f'<a href="{escape(run["link"])}">{text}</a>'
        out.append(text)
    return "".join(out)


def epub_paragraph(item, ebook):
    content = epub_runs(item["runs"])
    classes = []
    styles = []
    level = item.get("heading", 0)
    if level:
        if (ebook and level == 1) or item["align"] == "center":
            classes.append("center")
        attribute = f' class="{" ".join(classes)}"' if classes else ""
        tag = f"h{min(level, 3)}"
        return f"<{tag}{attribute}>{content or '&#160;'}</{tag}>"
    if item.get("line") and item["line"] != 100:
        styles.append(f"line-height: {item['line'] / 100:g}")
    if item.get("list"):
        marker = item["list"]["label"] or "\u2022"
        styles.append(f"margin-left: {1.5 * item['list']['level']:g}em")
        style_attribute = f' style="{"; ".join(styles)}"'
        return f'<p class="li"{style_attribute}>{escape(marker)}&#160;{content or "&#160;"}</p>'
    if item.get("left"):
        styles.append(f"margin-left: {2 * item['left']:g}em")
    if ebook:
        # KDP: no space between paragraphs; indented unless centered or right-aligned.
        classes.append("book")
        if not item["indent"]:
            classes.append("noindent")
    if item["align"] in ("center", "right", "justify"):
        classes.append(item["align"])
    attribute = f' class="{" ".join(classes)}"' if classes else ""
    if styles:
        attribute += f' style="{"; ".join(styles)}"'
    return f"<p{attribute}>{content or '&#160;'}</p>"


def epub_table(item, ebook):
    border = item["border"] or "#000000"
    lines = ["<table>"]
    for row in item["rows"]:
        lines.append("<tr>")
        for cell in row:
            style = [f"border-color: {border}"]
            if cell["background"]:
                style.append(f"background-color: {cell['background']}")
            inside = "".join(
                epub_paragraph(dict(p, heading=0, indent=False), False) for p in cell["paragraphs"]
            ) or "<p>&#160;</p>"
            lines.append(f'<td style="{"; ".join(style)}">{inside}</td>')
        lines.append("</tr>")
    lines.append("</table>")
    return "\n".join(lines)


def split_chapters(book, fallback_title):
    chapters = []
    for item in book:
        starts_chapter = item["type"] == "paragraph" and item["heading"] == 1
        if starts_chapter or not chapters:
            title = "".join(run["text"] for run in item["runs"]).strip() if starts_chapter else ""
            chapters.append({"title": title or fallback_title, "items": []})
        chapters[-1]["items"].append(item)
    return chapters or [{"title": fallback_title, "items": []}]


def xhtml_page(title, body):
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        "<!DOCTYPE html>\n"
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
        'lang="en" xml:lang="en">\n'
        f"<head><meta charset=\"utf-8\"/><title>{escape(title)}</title>"
        '<link rel="stylesheet" type="text/css" href="style.css"/></head>\n'
        f"<body>\n{body}\n</body>\n</html>\n"
    )


def write_epub(book, file_name, title, author, ebook):
    chapters = split_chapters(book, title)
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    # Pictures become files inside the EPUB.
    pictures = {}
    image_types = {"image/png": "png", "image/jpeg": "jpg", "image/gif": "gif"}

    def gather(paragraph):
        for run in paragraph.get("runs", []):
            if not run.get("image"):
                continue
            kind, data = image_bytes(run["image"])
            if kind not in image_types:
                continue
            key = run["image"]
            if key not in pictures:
                pictures[key] = (f"images/picture{len(pictures) + 1}.{image_types[kind]}", kind, data)
            run["image_file"] = pictures[key][0]

    for item in book:
        if item["type"] == "paragraph":
            gather(item)
        elif item["type"] == "table":
            for row in item["rows"]:
                for cell in row:
                    for paragraph in cell["paragraphs"]:
                        gather(paragraph)

    def epub_item(item):
        if item["type"] == "rule":
            return "<hr/>"
        if item["type"] == "paragraph":
            return epub_paragraph(item, ebook)
        return epub_table(item, ebook)

    pages = []
    for number, chapter in enumerate(chapters, start=1):
        body = "\n".join(epub_item(item) for item in chapter["items"])
        pages.append((f"chapter{number:03}.xhtml", chapter["title"], xhtml_page(chapter["title"], body)))

    nav_items = "\n".join(
        f'<li><a href="{name}">{escape(chapter_title)}</a></li>' for name, chapter_title, _ in pages
    )
    nav = xhtml_page(
        "Contents",
        f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>\n{nav_items}\n</ol></nav>',
    )

    manifest = "\n".join(
        f'<item id="c{number}" href="{name}" media-type="application/xhtml+xml"/>'
        for number, (name, _, _) in enumerate(pages, start=1)
    )
    for number, (name, kind, _) in enumerate(pictures.values(), start=1):
        manifest += f'\n<item id="p{number}" href="{name}" media-type="{kind}"/>'

    spine = "\n".join(f'<itemref idref="c{number}"/>' for number in range(1, len(pages) + 1))
    creator = f"<dc:creator>{escape(author)}</dc:creator>\n" if author else ""
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="en">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="bookid">{book_id}</dc:identifier>
<dc:title>{escape(title)}</dc:title>
{creator}<dc:language>en</dc:language>
<meta property="dcterms:modified">{modified}</meta>
</metadata>
<manifest>
<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>
<item id="css" href="style.css" media-type="text/css"/>
{manifest}
</manifest>
<spine toc="ncx">
{spine}
</spine>
</package>
"""

    nav_points = "\n".join(
        f'<navPoint id="n{number}" playOrder="{number}"><navLabel><text>{escape(chapter_title)}</text>'
        f'</navLabel><content src="{name}"/></navPoint>'
        for number, (name, chapter_title, _) in enumerate(pages, start=1)
    )
    ncx = f"""<?xml version="1.0" encoding="utf-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
<head><meta name="dtb:uid" content="{book_id}"/></head>
<docTitle><text>{escape(title)}</text></docTitle>
<navMap>
{nav_points}
</navMap>
</ncx>
"""

    container = """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>
"""

    with zipfile.ZipFile(file_name, "w") as target:
        target.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        target.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        target.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        target.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        target.writestr("OEBPS/toc.ncx", ncx, compress_type=zipfile.ZIP_DEFLATED)
        target.writestr("OEBPS/style.css", EPUB_CSS, compress_type=zipfile.ZIP_DEFLATED)
        for name, _, page in pages:
            target.writestr(f"OEBPS/{name}", page, compress_type=zipfile.ZIP_DEFLATED)
        for name, _, data in pictures.values():
            target.writestr(f"OEBPS/{name}", data, compress_type=zipfile.ZIP_STORED)


# ----- Reading Word files --------------------------------------------------

def docx_made_here(file_name):
    """Returns (made by Tabbed Writer?, is an eBook?) for a Word file."""
    from docx import Document
    category = Document(file_name).core_properties.category or ""
    return category.startswith("Tabbed Writer"), category == "Tabbed Writer eBook"


def docx_to_html(file_name):
    """Turn a Word file into simple HTML. Heading 1 paragraphs become <h1>."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(file_name)

    def paragraph_html(paragraph):
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        tag = {"Heading 1": "h1", "Title": "h1", "Heading 2": "h2", "Heading 3": "h3"}.get(style_name, "p")
        if style_name.startswith(("List Bullet", "List Number")):
            tag = "li"
        alignment = paragraph.alignment
        align = ""
        if alignment is not None:
            name = str(alignment).split(".")[-1].split(" ")[0].lower()
            if name.startswith("center"):
                align = "center"
            elif name.startswith("right"):
                align = "right"
            elif name.startswith("justify"):
                align = "justify"
        styles = []
        if align:
            styles.append(f"text-align: {align}")
        if paragraph.paragraph_format.page_break_before:
            styles.append("page-break-before: always")
        pieces = []
        if hasattr(paragraph, "iter_inner_content"):
            inner = list(paragraph.iter_inner_content())
        else:
            inner = list(paragraph.runs)
        runs = []
        for part in inner:
            if hasattr(part, "address"):  # a link
                for linked in part.runs:
                    runs.append((linked, part.address))
            else:
                runs.append((part, None))
        if not runs and paragraph._p.xpath("./w:pPr/w:pBdr/w:bottom") and not paragraph.text.strip():
            return "<hr/>"
        for run, address in runs:
            for embed in run._r.xpath(".//a:blip/@r:embed"):
                try:
                    picture = document.part.related_parts[embed]
                    widths = run._r.xpath(".//wp:extent/@cx")
                    width = int(int(widths[0]) / 9525) if widths else 0
                    data = base64.b64encode(picture.blob).decode()
                    size = f' width="{width}"' if width else ""
                    pieces.append(f'<img src="data:{picture.content_type};base64,{data}"{size}/>')
                except Exception:
                    pass
            text = escape(run.text).replace("\n", "<br/>").replace("\t", "&#160;&#160;&#160;&#160;")
            if not text:
                continue
            run_styles = []
            if run.font.color is not None and run.font.color.type is not None and run.font.color.rgb is not None:
                run_styles.append(f"color: #{run.font.color.rgb}")
            if run.font.size:
                run_styles.append(f"font-size: {run.font.size.pt:g}pt")
            if run.underline:
                text = f"<u>{text}</u>"
            if run.italic:
                text = f"<i>{text}</i>"
            if run.bold:
                text = f"<b>{text}</b>"
            if run_styles:
                text = f'<span style="{"; ".join(run_styles)}">{text}</span>'
            if address:
                text = f'<a href="{escape(address)}">{text}</a>'
            pieces.append(text)
        style_attribute = f' style="{"; ".join(styles)}"' if styles else ""
        return f"<{tag}{style_attribute}>{''.join(pieces)}</{tag}>"

    def table_html(table):
        rows = []
        for row in table.rows:
            cells = "".join(
                "<td>" + "".join(paragraph_html(p) for p in cell.paragraphs) + "</td>"
                for cell in row.cells
            )
            rows.append(f"<tr>{cells}</tr>")
        return '<table border="1" cellspacing="0" cellpadding="4" width="100%">' + "".join(rows) + "</table>"

    parts = []
    open_list = None
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        wanted = None
        if tag == "p":
            paragraph = Paragraph(child, document)
            style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style_name.startswith("List Bullet"):
                wanted = "ul"
            elif style_name.startswith("List Number"):
                wanted = "ol"
        if open_list and wanted != open_list:
            parts.append(f"</{open_list}>")
            open_list = None
        if wanted and not open_list:
            parts.append(f"<{wanted}>")
            open_list = wanted
        if tag == "p":
            parts.append(paragraph_html(paragraph))
        elif tag == "tbl":
            parts.append(table_html(Table(child, document)))
    if open_list:
        parts.append(f"</{open_list}>")
    return "<html><body>" + "\n".join(parts) + "</body></html>"


# ----- Turning the document into a "book" list ------------------------------

def brush_color(text_format, property_id):
    if text_format.hasProperty(property_id):
        brush = text_format.brushProperty(property_id)
        if brush.style() != Qt.BrushStyle.NoBrush:
            return brush.color().name()
    return None


def alignment_name(block_format):
    alignment = block_format.alignment()
    if alignment & Qt.AlignmentFlag.AlignHCenter:
        return "center"
    if alignment & Qt.AlignmentFlag.AlignRight:
        return "right"
    if alignment & Qt.AlignmentFlag.AlignJustify:
        return "justify"
    return "left"


def enum_number(value):
    """Qt sometimes hands back a plain number and sometimes a named value; compare as numbers."""
    return getattr(value, "value", value)


def block_to_paragraph(block):
    block_format = block.blockFormat()
    encoded = block.text().encode("utf-16-le")
    runs = []
    for part in block.textFormats():
        piece = encoded[2 * part.start: 2 * (part.start + part.length)].decode("utf-16-le", errors="replace")
        if not piece:
            continue
        text_format = part.format
        font = text_format.font()
        family = font.family()
        if text_format.isImageFormat():
            image_format = text_format.toImageFormat()
            if image_format.name().startswith("data:image/"):
                runs.append({
                    "text": "", "image": image_format.name(),
                    "width": image_format.width(), "height": image_format.height(),
                    "bold": False, "italic": False, "underline": False, "color": None,
                    "highlight": None, "font": None, "size": 0, "link": None,
                })
            continue
        vertical = text_format.verticalAlignment()
        runs.append({
            "link": text_format.anchorHref() if text_format.isAnchor() and text_format.anchorHref() else None,
            "text": piece,
            "strike": font.strikeOut(),
            "vertical": (
                "super" if vertical == QTextCharFormat.VerticalAlignment.AlignSuperScript
                else "sub" if vertical == QTextCharFormat.VerticalAlignment.AlignSubScript
                else None
            ),
            "bold": font.bold(),
            "italic": font.italic(),
            "underline": text_format.fontUnderline(),
            "color": brush_color(text_format, QTextFormat.Property.ForegroundBrush),
            "highlight": brush_color(text_format, QTextFormat.Property.BackgroundBrush),
            "font": family if family in FONT_LABELS and FONT_LABELS[family] != "Serif" else None,
            "size": text_format.fontPointSize() if text_format.fontPointSize() > 0 else 0,
        })
    text_list = block.textList()
    list_info = None
    if text_list is not None:
        list_format = text_list.format()
        bullet_styles = (
            QTextListFormat.Style.ListDisc,
            QTextListFormat.Style.ListCircle,
            QTextListFormat.Style.ListSquare,
        )
        is_bullet = list_format.style() in bullet_styles
        list_info = {
            "kind": "bullet" if is_bullet else "number",
            "level": max(1, list_format.indent()),
            "label": "" if is_bullet else text_list.itemText(block),
        }
    line = 100
    if enum_number(block_format.lineHeightType()) == enum_number(QTextBlockFormat.LineHeightTypes.ProportionalHeight):
        line = int(round(block_format.lineHeight())) or 100
    return {
        "type": "paragraph",
        "heading": block_format.headingLevel() if 0 < block_format.headingLevel() <= 3 else 0,
        "list": list_info,
        "left": block_format.indent(),
        "line": line,
        "align": alignment_name(block_format),
        "indent": block_format.textIndent() > 0,
        "page_break": bool(block_format.pageBreakPolicy() & QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore),
        "runs": runs,
    }


def table_to_item(document, table):
    rows = []
    for row in range(table.rows()):
        cells = []
        for column in range(table.columns()):
            cell = table.cellAt(row, column)
            paragraphs = []
            block = document.findBlock(cell.firstPosition())
            while block.isValid() and block.position() <= cell.lastPosition():
                paragraphs.append(block_to_paragraph(block))
                block = block.next()
            cells.append({
                "background": brush_color(cell.format(), QTextFormat.Property.BackgroundBrush),
                "paragraphs": paragraphs,
            })
        rows.append(cells)
    return {
        "type": "table",
        "border": brush_color(table.format(), QTextFormat.Property.FrameBorderBrush),
        "rows": rows,
    }


def document_to_book(document):
    book = []
    block = document.begin()
    while block.isValid():
        table = QTextCursor(block).currentTable()
        if table is not None:
            book.append(table_to_item(document, table))
            end = table.lastPosition()
            while block.isValid() and block.position() <= end:
                block = block.next()
            continue
        if block.blockFormat().hasProperty(QTextFormat.Property.BlockTrailingHorizontalRulerWidth):
            book.append({"type": "rule"})
        else:
            book.append(block_to_paragraph(block))
        block = block.next()
    # Leave out empty paragraphs at the very end.
    while book and book[-1]["type"] == "paragraph" and not "".join(r["text"] for r in book[-1]["runs"]).strip():
        book.pop()
    return book


# ----- PDF (made by Qt itself) ----------------------------------------------

def current_page_layout():
    width, height = TRIM_SIZES.get(page_settings["trim"], TRIM_SIZES[STANDARD_PAGE["trim"]])
    return QPageLayout(
        QPageSize(QSizeF(width, height), QPageSize.Unit.Inch, "", QPageSize.SizeMatchPolicy.ExactMatch),
        QPageLayout.Orientation.Portrait,
        QMarginsF(page_settings["inside"], page_settings["top"], page_settings["outside"], page_settings["bottom"]),
        QPageLayout.Unit.Inch,
    )


def paint_pages(copy, device):
    """Draw the document page by page, adding page numbers if they're turned on."""
    text_width, text_height = page_text_size()
    copy.setDocumentMargin(0)
    copy.setPageSize(QSizeF(text_width, text_height))
    scale = device.logicalDpiX() / SCREEN_DPI
    painter = QPainter()
    if not painter.begin(device):
        raise RuntimeError("Printing couldn't start.")
    number_font = QFont(DEFAULT_FONT)
    number_font.setPixelSize(13)
    try:
        for index in range(copy.pageCount()):
            if index:
                device.newPage()
            painter.save()
            painter.scale(scale, scale)
            painter.translate(0, -index * text_height)
            copy.drawContents(painter, QRectF(0, index * text_height, text_width, text_height))
            painter.restore()
            if page_settings.get("page_numbers"):
                painter.save()
                painter.scale(scale, scale)
                painter.setFont(number_font)
                below = text_height + page_settings["bottom"] * SCREEN_DPI * 0.35
                painter.drawText(QRectF(0, below, text_width, 24), Qt.AlignmentFlag.AlignHCenter, str(index + 1))
                painter.restore()
    finally:
        painter.end()


def write_pdf(copy, file_name, title):
    writer = QPdfWriter(file_name)
    writer.setTitle(title)
    writer.setCreator("Tabbed Writer")
    writer.setPageLayout(current_page_layout())
    paint_pages(copy, writer)


def print_document():
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setPageLayout(current_page_layout())
    printer.setDocName(DOCUMENT_FILE.stem if DOCUMENT_FILE else "Untitled")
    dialog = QPrintDialog(printer, window)
    dialog.setWindowTitle("Print")
    if not dialog.exec():
        return
    QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
    try:
        paint_pages(build_book_copy(), printer)
    except Exception as problem:
        QMessageBox.warning(window, "Print failed", str(problem))
    finally:
        QApplication.restoreOverrideCursor()


def word_support_missing():
    try:
        import docx  # noqa: F401
        return False
    except ImportError:
        return True


WORD_MISSING_MESSAGE = (
    "Word (.docx) files need the free python-docx add-on, which isn't installed.\n\n"
    "To add it, run this in a terminal:\n"
    "  ~/.local/share/tabbed-writer/.venv/bin/pip install python-docx"
)


# ---------------------------------------------------------------------------
# Saving copies (.docx, .epub, .pdf, .odt, .txt)
# ---------------------------------------------------------------------------

def is_centered_or_right(block_format):
    alignment = block_format.alignment()
    return bool(alignment & Qt.AlignmentFlag.AlignHCenter) or bool(alignment & Qt.AlignmentFlag.AlignRight)


def apply_book_paragraphs(document):
    # KDP: 0.2 inch first-line indent, no space before or after paragraphs.
    # Centered or right-aligned lines (like scene breaks) and tables are left alone.
    block = document.begin()
    while block.isValid():
        block_format = block.blockFormat()
        cursor = QTextCursor(block)
        if (
            block_format.headingLevel() == 0
            and block.textList() is None
            and cursor.currentTable() is None
            and block.text().strip()
            and not is_centered_or_right(block_format)
        ):
            book_paragraph = QTextBlockFormat()
            book_paragraph.setTextIndent(BOOK_INDENT)
            book_paragraph.setTopMargin(0)
            book_paragraph.setBottomMargin(0)
            cursor.mergeBlockFormat(book_paragraph)
        block = block.next()


def build_book_copy():
    current_snapshot()
    copy = QTextDocument()
    copy.setDefaultFont(QFont(DEFAULT_FONT, DEFAULT_SIZE))
    cursor = QTextCursor(copy)
    use_headings = sections.topLevelItemCount() > 1

    heading_text = QTextCharFormat()
    heading_text.setFontPointSize(18)
    heading_text.setFontWeight(QFont.Weight.Bold)

    ebook = document_kind == "ebook"
    for index in range(sections.topLevelItemCount()):
        item = sections.topLevelItem(index)
        if use_headings:
            heading_block = QTextBlockFormat()
            heading_block.setHeadingLevel(1)
            heading_block.setBottomMargin(24)
            if ebook:
                # KDP: each chapter title is a centered Heading 1 on a new page.
                heading_block.setAlignment(Qt.AlignmentFlag.AlignHCenter)
            if index == 0:
                cursor.setBlockFormat(heading_block)
            else:
                if ebook:
                    heading_block.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
                cursor.insertBlock(heading_block)
            cursor.insertText(item.text(0), heading_text)
            cursor.insertBlock(QTextBlockFormat(), QTextCharFormat())
        register_images(copy, item.data(0, CONTENT_ROLE))
        cursor.insertFragment(QTextDocumentFragment.fromHtml(item.data(0, CONTENT_ROLE) or ""))

    if ebook:
        apply_book_paragraphs(copy)
    return copy


def write_copy(file_name, snapshot=None, quiet=False):
    target = Path(file_name)
    if snapshot is None:
        snapshot = current_snapshot()
    copy = build_book_copy()

    if target.suffix == ".txt":
        writer = QTextDocumentWriter(str(target))
        writer.setFormat(b"plaintext")
        if not writer.write(copy):
            QMessageBox.warning(window, "Save failed", "The file could not be written.")
            return False
        window.statusBar().showMessage(f"Saved a copy as {target.name}", 4000)
        return True

    if target.suffix == ".docx" and word_support_missing():
        QMessageBox.warning(window, "Save failed", WORD_MISSING_MESSAGE)
        return False

    window.statusBar().showMessage(f"Saving {target.name}…")
    QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
    ebook = document_kind == "ebook"
    try:
        if target.suffix == ".pdf":
            write_pdf(copy, str(target), target.stem)
        elif target.suffix == ".docx":
            write_docx(document_to_book(copy), str(target), page_settings, TRIM_SIZES, target.stem, ebook)
        elif target.suffix == ".epub":
            write_epub(document_to_book(copy), str(target), target.stem, "", ebook)
        else:
            with tempfile.TemporaryDirectory() as work:
                odt = Path(work) / (target.stem + ".odt")
                writer = QTextDocumentWriter(str(odt))
                writer.setFormat(b"odf")
                if not writer.write(copy):
                    raise RuntimeError("The file could not be written.")
                finish_odt(odt, page_settings, target.stem)
                tuck_tab_data(odt, snapshot)
                shutil.copyfile(odt, target)
    except Exception as problem:
        QApplication.restoreOverrideCursor()
        window.statusBar().clearMessage()
        QMessageBox.warning(window, "Save failed", str(problem))
        return False
    QApplication.restoreOverrideCursor()
    if not quiet:
        window.statusBar().showMessage(f"Saved a copy as {target.name}", 4000)
    return True


def save_as_document():
    global DOCUMENT_FILE, FROM_OTHER_PROGRAM
    if DOCUMENT_FILE is None:
        start = WRITING_FOLDER / (Path(SUGGESTED_NAME).stem + default_ending())
    else:
        start = DOCUMENT_FILE
    start_ending = start.suffix.lower()
    start_filter = next(
        (label for label, ending in SAVE_TYPES.items() if ending == start_ending),
        next(iter(SAVE_TYPES)),
    )
    file_name, chosen = QFileDialog.getSaveFileName(
        window,
        "Save as",
        str(start),
        ";;".join(SAVE_TYPES),
        start_filter,
    )
    if not file_name:
        return

    ending = SAVE_TYPES.get(chosen) or Path(file_name).suffix.lower() or default_ending()
    if ending not in SAVE_TYPES.values():
        ending = default_ending()
    if not file_name.lower().endswith(ending):
        file_name += ending

    if ending in WORKING_FORMATS:
        if ending == ".docx" and sections.topLevelItemCount() > 1:
            QMessageBox.information(
                window,
                "Saving as a Word file",
                "Word files can't hold tabs, so each tab is saved as a heading. "
                "When you open the file in Tabbed Writer again, the headings turn back into tabs.",
            )
        previous_file, previous_flag = DOCUMENT_FILE, FROM_OTHER_PROGRAM
        DOCUMENT_FILE, FROM_OTHER_PROGRAM = Path(file_name), False
        save_document()
        if current_snapshot() != saved_snapshot:
            DOCUMENT_FILE, FROM_OTHER_PROGRAM = previous_file, previous_flag
            return
        window.setWindowTitle(f"{DOCUMENT_FILE.name} - Tabbed Writer")
        return

    if sections.topLevelItemCount() > 1:
        QMessageBox.information(
            window,
            "Saving a copy",
            f"{ending} files can't keep tabs, so each tab's name becomes a heading in the copy"
            + (" (a centered chapter title on a new page)." if document_kind == "ebook" else ".")
            + "\n\nYour tabs stay as they are here.",
        )
    write_copy(file_name)


# ---------------------------------------------------------------------------
# Opening .twp, .odt, .docx and .txt
# ---------------------------------------------------------------------------

def html_from_plain_text(text):
    holder = QTextDocument()
    holder.setPlainText(text)
    return holder.toHtml()


def html_from_office_file(path):
    with tempfile.TemporaryDirectory() as work:
        libreoffice_convert(path, "html", work)
        result = Path(work) / (path.stem + ".html")
        if not result.exists():
            raise RuntimeError("LibreOffice couldn't read that file.")
        return result.read_text(encoding="utf-8", errors="replace")


def tabs_from_html(html):
    # Each Heading 1 (a chapter title) becomes its own tab.
    document = QTextDocument()
    document.setHtml(html)

    headings = []
    block = document.begin()
    while block.isValid():
        if (
            block.blockFormat().headingLevel() == 1
            and block.text().strip()
            and QTextCursor(block).currentTable() is None
        ):
            headings.append(block)
        block = block.next()

    if not headings:
        return [{"title": "Tab 1", "text": html}]

    def piece(start, end):
        if start >= end:
            return ""
        cursor = QTextCursor(document)
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
        if not cursor.selectedText().strip():
            return ""
        return QTextDocumentFragment(cursor).toHtml()

    tabs = []
    opening = piece(0, headings[0].position() - 1)
    if opening:
        tabs.append({"title": "Front matter", "text": opening})
    last_position = document.characterCount() - 1
    for number, heading in enumerate(headings):
        start = heading.position() + heading.length()
        if number + 1 < len(headings):
            end = headings[number + 1].position() - 1
        else:
            end = last_position
        tabs.append({"title": heading.text().strip(), "text": piece(start, end)})
    return tabs


def open_document(file_name=None):
    global DOCUMENT_FILE, saved_snapshot, SUGGESTED_NAME, page_settings, FROM_OTHER_PROGRAM
    if not ok_to_discard():
        return
    if not isinstance(file_name, str):
        file_name, _ = QFileDialog.getOpenFileName(
            window,
            "Open",
            str(WRITING_FOLDER),
            "All supported (*.twp *.docx *.odt *.txt);;Tabbed Writer document (*.twp);;"
            "Word document (*.docx);;OpenDocument text (*.odt);;Plain text (*.txt)",
        )
    if not file_name:
        return
    path = Path(file_name)
    if path.parent == RECOVERY_FOLDER:
        try:
            name = recover_from(path)
        except (OSError, ValueError) as problem:
            QMessageBox.warning(window, "Couldn't recover", str(problem))
            return
        window.setWindowTitle(f"{name} (recovered) - Tabbed Writer")
        saved_snapshot = ""  # counts as unsaved until you save
        return
    if not load_file(path):
        return
    window.setWindowTitle(f"{path.name} - Tabbed Writer")
    saved_snapshot = current_snapshot()
    clear_recovery()
    if DOCUMENT_FILE is not None:
        remember_recent(DOCUMENT_FILE)


def load_file(path):
    """Fill this window from a file. Returns True if it worked."""
    global DOCUMENT_FILE, SUGGESTED_NAME, page_settings, FROM_OTHER_PROGRAM
    kind = path.suffix.lower()
    try:
        stored = read_tab_data(path) if kind == ".odt" else None
        if kind == ".twp":
            load_document(path)
            DOCUMENT_FILE, FROM_OTHER_PROGRAM = path, False
        elif stored is not None:
            # An .odt saved by Tabbed Writer: bring the tabs back exactly.
            apply_saved_data(stored)
            DOCUMENT_FILE, FROM_OTHER_PROGRAM = path, False
        else:
            made_here, ebook = False, False
            if kind == ".txt":
                tabs = [{"title": "Tab 1", "text": html_from_plain_text(
                    path.read_text(encoding="utf-8", errors="replace"))}]
            elif kind == ".docx":
                if word_support_missing():
                    raise RuntimeError(WORD_MISSING_MESSAGE)
                made_here, ebook = docx_made_here(str(path))
                tabs = tabs_from_html(docx_to_html(str(path)))
            else:
                tabs = tabs_from_html(html_from_office_file(path))
            page_settings = dict(DEFAULT_PAGE if ebook else STANDARD_PAGE)
            set_document_kind("ebook" if ebook else "standard")
            apply_page_look()
            fill_tabs(tabs)
            if kind == ".txt":
                DOCUMENT_FILE, FROM_OTHER_PROGRAM = None, False
                SUGGESTED_NAME = path.stem + default_ending()
            else:
                DOCUMENT_FILE, FROM_OTHER_PROGRAM = path, not made_here
    except Exception as problem:
        QMessageBox.warning(window, "Couldn't open file", str(problem))
        return False
    return True


def new_document(kind="standard"):
    global DOCUMENT_FILE, saved_snapshot, SUGGESTED_NAME, page_settings, FROM_OTHER_PROGRAM
    if not ok_to_discard():
        return
    DOCUMENT_FILE = None
    FROM_OTHER_PROGRAM = False
    SUGGESTED_NAME = "Untitled" + (".twp" if kind == "ebook" else ".odt")
    page_settings = dict(DEFAULT_PAGE if kind == "ebook" else STANDARD_PAGE)
    set_document_kind(kind)
    apply_page_look()
    first_title = "Chapter 1" if kind == "ebook" else "Tab 1"
    fill_tabs([{"title": first_title, "text": ""}])
    window.setWindowTitle("Untitled - Tabbed Writer")
    saved_snapshot = current_snapshot()
    clear_recovery()


def toggle_ebook_formatting(turned_on):
    global document_kind
    document_kind = "ebook" if turned_on else "standard"
    use_book_paragraph_if_empty()
    update_counts()
    if turned_on:
        window.statusBar().showMessage(
            "eBook formatting on: saved copies follow the KDP and B&N rules", 5000
        )


# ---------------------------------------------------------------------------
# Page setup (for paperbacks)
# ---------------------------------------------------------------------------

def page_setup():
    global page_settings
    dialog = QDialog(window)
    dialog.setWindowTitle("Page setup (paperback)")
    form = QFormLayout(dialog)

    trim_box = QComboBox()
    trim_box.addItems(list(TRIM_SIZES))
    trim_box.setCurrentText(page_settings["trim"])
    form.addRow("Trim size:", trim_box)

    margin_boxes = {}
    for key, label in [("top", "Top"), ("bottom", "Bottom"),
                       ("inside", "Inside (gutter)"), ("outside", "Outside")]:
        box = QDoubleSpinBox()
        box.setRange(0.25, 2.0)
        box.setSingleStep(0.125)
        box.setDecimals(3)
        box.setSuffix(" in")
        box.setValue(float(page_settings[key]))
        margin_boxes[key] = box
        form.addRow(f"{label}:", box)

    form.addRow(QLabel(
        "KDP minimum inside margin by page count:\n"
        "  24–150 pages: 0.375 in      151–300: 0.5 in\n"
        "  301–500: 0.625 in      501–700: 0.75 in      701–828: 0.875 in\n"
        "Outside margin: at least 0.25 in.\n\n"
        "These settings set the page on screen and in PDF, Word and .odt copies.\n"
        "eBook readers ignore them, since readers choose their own layout."
    ))

    numbers_box = QCheckBox("Page numbers at the bottom (PDF, printing and Word)")
    numbers_box.setChecked(bool(page_settings.get("page_numbers")))
    form.addRow(numbers_box)
    form.addRow(QLabel("KDP asks for no page numbers in eBooks; they're for paperbacks and printing."))

    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
    buttons.accepted.connect(dialog.accept)
    buttons.rejected.connect(dialog.reject)
    form.addRow(buttons)

    if not dialog.exec():
        return
    page_settings = {"trim": trim_box.currentText(), "page_numbers": numbers_box.isChecked()}
    for key, box in margin_boxes.items():
        page_settings[key] = round(box.value(), 3)
    apply_page_look()
    update_counts()


# ---------------------------------------------------------------------------
# Insert: breaks and tables
# ---------------------------------------------------------------------------

def insert_page_break():
    cursor = editor.textCursor()
    break_format = cursor.blockFormat()
    break_format.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
    break_format.setHeadingLevel(0)
    cursor.insertBlock(break_format)
    editor.setTextCursor(cursor)
    window.statusBar().showMessage("Page break added (it shows in saved copies)", 4000)


def insert_line_break():
    editor.textCursor().insertText(" ")


def insert_table():
    dialog = QDialog(window)
    dialog.setWindowTitle("Insert table")
    form = QFormLayout(dialog)
    rows_box = QSpinBox()
    rows_box.setRange(1, 50)
    rows_box.setValue(3)
    columns_box = QSpinBox()
    columns_box.setRange(1, 10)
    columns_box.setValue(3)
    form.addRow("Rows:", rows_box)
    form.addRow("Columns:", columns_box)
    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
    buttons.accepted.connect(dialog.accept)
    buttons.rejected.connect(dialog.reject)
    form.addRow(buttons)
    if not dialog.exec():
        return

    table_format = QTextTableFormat()
    table_format.setBorder(1)
    table_format.setBorderStyle(QTextFrameFormat.BorderStyle.BorderStyle_Solid)
    table_format.setBorderCollapse(True)
    table_format.setCellPadding(4)
    table_format.setCellSpacing(0)
    table_format.setWidth(QTextLength(QTextLength.Type.PercentageLength, 100))
    editor.textCursor().insertTable(rows_box.value(), columns_box.value(), table_format)
    editor.setFocus()


def current_table_cell():
    cursor = editor.textCursor()
    table = cursor.currentTable()
    if table is None:
        return None, None
    return table, table.cellAt(cursor)


def add_table_row(offset):
    table, cell = current_table_cell()
    if table is not None:
        table.insertRows(cell.row() + offset, 1)


def add_table_column(offset):
    table, cell = current_table_cell()
    if table is not None:
        table.insertColumns(cell.column() + offset, 1)


def delete_table():
    table, _ = current_table_cell()
    if table is not None:
        table.removeRows(0, table.rows())


def delete_table_row():
    table, cell = current_table_cell()
    if table is None:
        return
    if table.rows() == 1:
        delete_table()
    else:
        table.removeRows(cell.row(), 1)


def delete_table_column():
    table, cell = current_table_cell()
    if table is None:
        return
    if table.columns() == 1:
        delete_table()
    else:
        table.removeColumns(cell.column(), 1)


def choose_color(title, start_color):
    color = QColorDialog.getColor(
        start_color, window, title, QColorDialog.ColorDialogOption.DontUseNativeDialog
    )
    return color if color.isValid() else None


def selected_table_cells():
    cursor = editor.textCursor()
    table = cursor.currentTable()
    if table is None:
        return None, []
    if cursor.hasComplexSelection():
        first_row, row_count, first_column, column_count = cursor.selectedTableCells()
    else:
        cell = table.cellAt(cursor)
        first_row, row_count, first_column, column_count = cell.row(), 1, cell.column(), 1
    cells = []
    for row in range(first_row, first_row + row_count):
        for column in range(first_column, first_column + column_count):
            cells.append(table.cellAt(row, column))
    return table, cells


def choose_cell_color():
    table, cells = selected_table_cells()
    if not cells:
        return
    color = choose_color("Cell color", QColor("#fff2cc"))
    if color is None:
        return
    for cell in cells:
        cell_format = cell.format()
        cell_format.setBackground(color)
        cell.setFormat(cell_format)


def clear_cell_color():
    table, cells = selected_table_cells()
    for cell in cells:
        cell_format = cell.format()
        cell_format.clearBackground()
        cell.setFormat(cell_format)


def choose_border_color():
    table, _ = selected_table_cells()
    if table is None:
        return
    table_format = table.format()
    color = choose_color("Border color", table_format.borderBrush().color())
    if color is None:
        return
    table_format.setBorderBrush(QBrush(color))
    table.setFormat(table_format)


def show_editor_menu(position):
    if not editor.textCursor().hasSelection():
        editor.setTextCursor(editor.cursorForPosition(position))
    menu = editor.createStandardContextMenu(position)
    misspelled = misspelled_word_at(editor.cursorForPosition(position))
    if misspelled is not None:
        start, end, word = misspelled
        first = menu.actions()[0] if menu.actions() else None
        spelling = []
        choices = suggestions_for(word)
        for choice in choices:
            action = QAction(choice, menu)
            bold = action.font()
            bold.setBold(True)
            action.setFont(bold)
            action.triggered.connect(lambda checked=False, c=choice: replace_range(start, end, c))
            spelling.append(action)
        if not choices:
            none_found = QAction("(No suggestions)", menu)
            none_found.setEnabled(False)
            spelling.append(none_found)
        add_word = QAction(f"Add “{word}” to dictionary", menu)
        add_word.triggered.connect(lambda checked=False: add_to_dictionary(word))
        ignore = QAction("Ignore", menu)
        ignore.triggered.connect(lambda checked=False: ignore_word(word))
        spelling += [add_word, ignore]
        for action in spelling:
            menu.insertAction(first, action)
        if first is not None:
            menu.insertSeparator(first)
    clicked = editor.cursorForPosition(position)
    picked = image_cursor_at(clicked)
    if picked is not None:
        menu.addSeparator()
        size_menu = menu.addMenu("Picture size")
        for label, fraction in (("Small (25%)", 0.25), ("Medium (50%)", 0.5), ("Large (75%)", 0.75), ("Full width", 1.0)):
            size_menu.addAction(label, lambda checked=False, f=fraction, c=picked: resize_image(c, f))
    link = link_range_at(clicked)
    if link is not None:
        menu.addSeparator()
        menu.addAction("Open link", lambda checked=False, a=link[2]: QDesktopServices.openUrl(QUrl(a)))
        menu.addAction("Edit link…", insert_link)
        menu.addAction("Remove link", lambda checked=False, s=link[0], e=link[1]: remove_link(s, e))
    if editor.textCursor().currentTable() is not None:
        menu.addSeparator()
        table_menu = menu.addMenu("Table")
        table_menu.addAction("Insert row above", lambda: add_table_row(0))
        table_menu.addAction("Insert row below", lambda: add_table_row(1))
        table_menu.addAction("Insert column left", lambda: add_table_column(0))
        table_menu.addAction("Insert column right", lambda: add_table_column(1))
        table_menu.addSeparator()
        table_menu.addAction("Delete row", delete_table_row)
        table_menu.addAction("Delete column", delete_table_column)
        table_menu.addAction("Delete table", delete_table)
        table_menu.addSeparator()
        table_menu.addAction("Cell color…", choose_cell_color)
        table_menu.addAction("Remove cell color", clear_cell_color)
        table_menu.addAction("Border color…", choose_border_color)
    menu.exec(editor.viewport().mapToGlobal(position))



# ---------------------------------------------------------------------------
# Spell check (offline, using the pyspellchecker add-on)
# ---------------------------------------------------------------------------

WORD_PATTERN = re.compile(r"[^\W\d_]+(?:['’][^\W\d_]+)*")
CONTRACTION_ENDINGS = ("n't", "'s", "'ll", "'re", "'ve", "'d", "'m")
ALWAYS_KNOWN = {"can't", "won't", "shan't", "ain't"}
SPELL_MISSING_MESSAGE = (
    "Spell check needs the free pyspellchecker add-on, which isn't installed.\n\n"
    "To add it, run this in a terminal:\n"
    "  ~/.local/share/tabbed-writer/.venv/bin/pip install pyspellchecker"
)

spell = None            # the dictionary, loaded shortly after the window opens
spell_cache = {}        # words already looked up
ignored_words = set()   # "Ignore" for this session only


def stored_words():
    words = app_settings.value("dictionary_words", [])
    if isinstance(words, str):
        words = [words]
    return [word for word in (words or []) if word]


def load_spell_checker():
    global spell
    try:
        from spellchecker import SpellChecker
    except ImportError:
        spell_action.setEnabled(False)
        spell_action.setToolTip("Needs the pyspellchecker add-on")
        return
    spell = SpellChecker(language="en")
    own_words = stored_words()
    if own_words:
        spell.word_frequency.load_words(own_words)
    spell_cache.clear()
    highlighter.rehighlight()


def word_is_known(word):
    key = word.lower().replace("’", "'")
    if key in spell_cache:
        return spell_cache[key]
    known = key in ignored_words or key in ALWAYS_KNOWN or bool(spell.known([key]))
    if not known:
        for ending in CONTRACTION_ENDINGS:
            if key.endswith(ending) and len(key) > len(ending) and spell.known([key[: -len(ending)]]):
                known = True
                break
    spell_cache[key] = known
    return known


def word_needs_checking(word):
    if len(word) < 2:
        return False
    if word.isupper():
        return False  # acronyms like NASA
    return True


def utf16_offsets(text):
    """Qt counts some characters (like emoji) as two; map Python positions to Qt's."""
    if all(ord(character) < 0x10000 for character in text):
        return None
    offsets = [0]
    for character in text:
        offsets.append(offsets[-1] + (2 if ord(character) >= 0x10000 else 1))
    return offsets


def misspelled_spans(text):
    if spell is None:
        return []
    offsets = utf16_offsets(text)
    spans = []
    for match in WORD_PATTERN.finditer(text):
        word = match.group()
        if word_needs_checking(word) and not word_is_known(word):
            start, end = match.start(), match.end()
            if offsets:
                start, end = offsets[start], offsets[end]
            spans.append((start, end, word))
    return spans


class SpellHighlighter(QSyntaxHighlighter):
    def __init__(self, document):
        super().__init__(document)
        self.enabled = True
        self.mark = QTextCharFormat()
        self.mark.setUnderlineStyle(QTextCharFormat.UnderlineStyle.SpellCheckUnderline)
        self.mark.setUnderlineColor(QColor("#d93025"))

    def highlightBlock(self, text):
        if not self.enabled:
            return
        for start, end, _ in misspelled_spans(text):
            self.setFormat(start, end - start, self.mark)


def misspelled_word_at(cursor):
    """Returns (start, end, word) for a misspelled word under the cursor, or None."""
    if spell is None or not highlighter.enabled:
        return None
    block = cursor.block()
    offset = cursor.positionInBlock()
    for start, end, word in misspelled_spans(block.text()):
        if start <= offset <= end:
            return block.position() + start, block.position() + end, word
    return None


def suggestions_for(word):
    key = word.lower().replace("’", "'")
    old_distance = spell.distance
    try:
        if len(key) > 9:
            spell.distance = 1
        found = spell.candidates(key) or set()
    except Exception:
        found = set()
    finally:
        spell.distance = old_distance
    found.discard(key)
    ordered = sorted(found, key=lambda candidate: -spell[candidate])[:5]
    if word[:1].isupper():
        ordered = [candidate[:1].upper() + candidate[1:] for candidate in ordered]
    return ordered


def replace_range(start, end, text):
    cursor = editor.textCursor()
    cursor.setPosition(start)
    cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
    cursor.insertText(text)


def add_to_dictionary(word):
    key = word.lower().replace("’", "'")
    words = stored_words()
    if key not in words:
        words.append(key)
        app_settings.setValue("dictionary_words", words)
    spell.word_frequency.load_words([key])
    spell_cache.pop(key, None)
    highlighter.rehighlight()


def ignore_word(word):
    key = word.lower().replace("’", "'")
    ignored_words.add(key)
    spell_cache.pop(key, None)
    highlighter.rehighlight()


def set_spell_check(turned_on):
    if turned_on and spell is None:
        try:
            import spellchecker  # noqa: F401
        except ImportError:
            QMessageBox.information(window, "Spell check", SPELL_MISSING_MESSAGE)
            spell_action.setChecked(False)
            return
    app_settings.setValue("spell_check", turned_on)
    highlighter.enabled = turned_on
    highlighter.rehighlight()


# ---------------------------------------------------------------------------
# Find and replace
# ---------------------------------------------------------------------------

def html_has(html, text, flags):
    holder = QTextDocument()
    holder.setHtml(html or "")
    return not holder.find(text, 0, flags).isNull()


def replace_in_document(document, text, replacement, flags):
    count = 0
    edit = QTextCursor(document)
    edit.beginEditBlock()
    found = document.find(text, 0, flags)
    while not found.isNull():
        found.insertText(replacement)
        count += 1
        found = document.find(text, found.position(), flags)
    edit.endEditBlock()
    return count


class FindReplaceDialog(QDialog):
    def __init__(self):
        super().__init__(window)
        self.setWindowTitle("Find and replace")
        form = QFormLayout(self)
        self.find_box = QLineEdit()
        self.replace_box = QLineEdit()
        form.addRow("Find:", self.find_box)
        form.addRow("Replace with:", self.replace_box)

        self.match_case = QCheckBox("Match case")
        self.whole_words = QCheckBox("Whole words only")
        self.all_tabs = QCheckBox("Search all tabs")
        self.all_tabs.setChecked(True)
        options = QHBoxLayout()
        for box in (self.match_case, self.whole_words, self.all_tabs):
            options.addWidget(box)
        form.addRow(options)

        self.status = QLabel()
        form.addRow(self.status)

        buttons = QHBoxLayout()
        for label, action in (
            ("Previous", lambda: self.find_next(backward=True)),
            ("Next", lambda: self.find_next()),
            ("Replace", self.replace_one),
            ("Replace all", self.replace_all),
            ("Close", self.close),
        ):
            button = QPushButton(label)
            button.setAutoDefault(False)
            button.clicked.connect(action)
            buttons.addWidget(button)
        form.addRow(buttons)

        self.find_box.returnPressed.connect(lambda: self.find_next())
        self.replace_box.returnPressed.connect(self.replace_one)

    def flags(self, backward=False):
        flags = QTextDocument.FindFlag(0)
        if self.match_case.isChecked():
            flags |= QTextDocument.FindFlag.FindCaseSensitively
        if self.whole_words.isChecked():
            flags |= QTextDocument.FindFlag.FindWholeWords
        if backward:
            flags |= QTextDocument.FindFlag.FindBackward
        return flags

    def find_next(self, backward=False):
        text = self.find_box.text()
        if not text:
            return False
        flags = self.flags(backward)
        if editor.find(text, flags):
            self.status.clear()
            return True
        count = sections.topLevelItemCount()
        start = sections.indexOfTopLevelItem(sections.currentItem())
        if self.all_tabs.isChecked() and count > 1:
            order = [(start + (-step if backward else step)) % count for step in range(1, count + 1)]
        else:
            order = [start]
        search_flags = self.flags()
        for index in order:
            item = sections.topLevelItem(index)
            if item is not sections.currentItem():
                if not html_has(item.data(0, CONTENT_ROLE), text, search_flags):
                    continue
                sections.setCurrentItem(item)
            cursor = editor.textCursor()
            cursor.movePosition(
                QTextCursor.MoveOperation.End if backward else QTextCursor.MoveOperation.Start
            )
            editor.setTextCursor(cursor)
            if editor.find(text, flags):
                self.status.clear()
                return True
        self.status.setText("No matches found.")
        return False

    def selection_matches(self):
        selected = editor.textCursor().selectedText()
        wanted = self.find_box.text()
        if not selected or not wanted:
            return False
        if self.match_case.isChecked():
            return selected == wanted
        return selected.lower() == wanted.lower()

    def replace_one(self):
        if self.selection_matches():
            editor.textCursor().insertText(self.replace_box.text())
        self.find_next()

    def replace_all(self):
        text = self.find_box.text()
        if not text:
            return
        replacement = self.replace_box.text()
        flags = self.flags()
        current = sections.currentItem()
        if self.all_tabs.isChecked():
            items = [sections.topLevelItem(index) for index in range(sections.topLevelItemCount())]
        else:
            items = [current]
        total = 0
        for item in items:
            if item is current:
                total += replace_in_document(editor.document(), text, replacement, flags)
                continue
            holder = QTextDocument()
            holder.setHtml(item.data(0, CONTENT_ROLE) or "")
            changed = replace_in_document(holder, text, replacement, flags)
            if changed:
                item.setData(0, CONTENT_ROLE, holder.toHtml())
                item.setData(0, WORDS_ROLE, count_words(holder.toPlainText()))
                total += changed
        self.status.setText(f"Replaced {total:,} " + ("match." if total == 1 else "matches."))
        update_counts_soon()


find_dialog = None


def show_find(with_replace):
    global find_dialog
    if find_dialog is None:
        find_dialog = FindReplaceDialog()
    selected = editor.textCursor().selectedText()
    if selected and " " not in selected:
        find_dialog.find_box.setText(selected)
    find_dialog.setWindowTitle("Find and replace" if with_replace else "Find")
    find_dialog.status.clear()
    find_dialog.show()
    find_dialog.raise_()
    find_dialog.activateWindow()
    if with_replace and find_dialog.find_box.text():
        find_dialog.replace_box.setFocus()
    else:
        find_dialog.find_box.setFocus()
        find_dialog.find_box.selectAll()


# ---------------------------------------------------------------------------
# Lists, indents, spacing and paragraph styles
# ---------------------------------------------------------------------------

BULLET_STYLES = (
    QTextListFormat.Style.ListDisc,
    QTextListFormat.Style.ListCircle,
    QTextListFormat.Style.ListSquare,
)
NUMBER_STYLES = (
    QTextListFormat.Style.ListDecimal,
    QTextListFormat.Style.ListLowerAlpha,
    QTextListFormat.Style.ListLowerRoman,
)


def selected_blocks(cursor=None):
    cursor = cursor or editor.textCursor()
    document = editor.document()
    block = document.findBlock(cursor.selectionStart())
    last = document.findBlock(cursor.selectionEnd())
    blocks = []
    while block.isValid():
        blocks.append(block)
        if block == last:
            break
        block = block.next()
    return blocks


def list_styles_for(style):
    return BULLET_STYLES if style in BULLET_STYLES else NUMBER_STYLES


def take_out_of_list(block):
    text_list = block.textList()
    if text_list is not None:
        text_list.remove(block)
    block_format = block.blockFormat()
    block_format.setIndent(0)
    QTextCursor(block).setBlockFormat(block_format)


def toggle_list(style):
    cursor = editor.textCursor()
    current = cursor.currentList()
    cursor.beginEditBlock()
    if current is not None and current.format().style() in list_styles_for(style):
        for block in selected_blocks(cursor):
            take_out_of_list(block)
    elif current is not None:
        list_format = current.format()
        styles = list_styles_for(style)
        list_format.setStyle(styles[min(max(list_format.indent(), 1), 3) - 1])
        current.setFormat(list_format)
    else:
        list_format = QTextListFormat()
        list_format.setStyle(style)
        list_format.setIndent(1)
        cursor.createList(list_format)
    cursor.endEditBlock()
    editor.setFocus()


def neighbour_list(block, level, styles):
    previous = block.previous()
    while previous.isValid() and previous.textList() is not None:
        text_list = previous.textList()
        if text_list.format().indent() == level and text_list.format().style() in styles:
            return text_list
        previous = previous.previous()
    return None


def change_indent(step):
    cursor = editor.textCursor()
    cursor.beginEditBlock()
    for block in selected_blocks(cursor):
        text_list = block.textList()
        if text_list is None:
            block_format = block.blockFormat()
            block_format.setIndent(max(0, block_format.indent() + step))
            QTextCursor(block).setBlockFormat(block_format)
            continue
        list_format = QTextListFormat(text_list.format())
        level = list_format.indent() + step
        if level < 1:
            take_out_of_list(block)
            continue
        level = min(level, 6)
        styles = list_styles_for(list_format.style())
        list_format.setIndent(level)
        list_format.setStyle(styles[(level - 1) % 3])
        target = neighbour_list(block, level, styles)
        text_list.remove(block)
        if target is not None:
            target.add(block)
        else:
            QTextCursor(block).createList(list_format)
    cursor.endEditBlock()
    editor.setFocus()


def set_line_spacing(percent):
    cursor = editor.textCursor()
    cursor.beginEditBlock()
    for block in selected_blocks(cursor):
        block_format = block.blockFormat()
        block_format.setLineHeight(
            float(percent), enum_number(QTextBlockFormat.LineHeightTypes.ProportionalHeight)
        )
        QTextCursor(block).setBlockFormat(block_format)
    cursor.endEditBlock()
    editor.setFocus()


# Like Google Docs: name -> (heading level, font size, color)
PARAGRAPH_STYLES = {
    "Normal text": (0, DEFAULT_SIZE, None),
    "Title": (0, 26, None),
    "Subtitle": (0, 15, "#666666"),
    "Heading 1": (1, 20, None),
    "Heading 2": (2, 16, None),
    "Heading 3": (3, 14, "#434343"),
}
STYLE_GRAYS = ("#666666", "#434343")


def whole_block_cursor(block):
    cursor = QTextCursor(block)
    cursor.movePosition(QTextCursor.MoveOperation.StartOfBlock)
    cursor.movePosition(QTextCursor.MoveOperation.EndOfBlock, QTextCursor.MoveMode.KeepAnchor)
    return cursor


def apply_paragraph_style(name):
    level, size, color = PARAGRAPH_STYLES[name]
    cursor = editor.textCursor()
    cursor.beginEditBlock()
    for block in selected_blocks(cursor):
        was_heading = block.blockFormat().headingLevel() > 0
        block_format = block.blockFormat()
        block_format.setHeadingLevel(level)
        if level:
            block_format.setTopMargin(12)
            block_format.setBottomMargin(6)
        QTextCursor(block).setBlockFormat(block_format)

        # Remove the gray used by Subtitle and Heading 3 before applying the new look.
        for part in block.textFormats():
            if part.format.foreground().color().name() in STYLE_GRAYS:
                piece = QTextCursor(block)
                piece.setPosition(block.position() + part.start)
                piece.setPosition(block.position() + part.start + part.length, QTextCursor.MoveMode.KeepAnchor)
                plain = part.format
                plain.clearForeground()
                piece.setCharFormat(plain)

        look = QTextCharFormat()
        look.setFontPointSize(size)
        if color:
            look.setForeground(QColor(color))
        if level:
            look.setFontWeight(QFont.Weight.Bold)
        elif was_heading:
            look.setFontWeight(QFont.Weight.Normal)
        whole = whole_block_cursor(block)
        whole.mergeCharFormat(look)
        whole.mergeBlockCharFormat(look)
    cursor.endEditBlock()
    editor.mergeCurrentCharFormat(look)
    editor.setFocus()


def current_style_name():
    level = editor.textCursor().block().blockFormat().headingLevel()
    if level in (1, 2, 3):
        return f"Heading {level}"
    size = round(editor.currentCharFormat().fontPointSize())
    if size == 26:
        return "Title"
    if size == 15:
        return "Subtitle"
    return "Normal text"


def update_style_box():
    style_box.blockSignals(True)
    style_box.setCurrentText(current_style_name())
    style_box.blockSignals(False)


# ---------------------------------------------------------------------------
# More text formatting
# ---------------------------------------------------------------------------

def toggle_strikethrough():
    look = QTextCharFormat()
    look.setFontStrikeOut(not editor.currentCharFormat().fontStrikeOut())
    editor.mergeCurrentCharFormat(look)


def toggle_vertical(alignment):
    look = QTextCharFormat()
    if editor.currentCharFormat().verticalAlignment() == alignment:
        look.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignNormal)
    else:
        look.setVerticalAlignment(alignment)
    editor.mergeCurrentCharFormat(look)


def clear_formatting():
    cursor = editor.textCursor()
    if cursor.hasSelection():
        cursor.setCharFormat(QTextCharFormat())
    else:
        editor.setCurrentCharFormat(QTextCharFormat())
    editor.setFocus()


def paste_without_formatting():
    text = QApplication.clipboard().text()
    if text:
        editor.insertPlainText(text)



# ---------------------------------------------------------------------------
# Pictures, links and lines
# ---------------------------------------------------------------------------

def insert_image():
    file_name, _ = QFileDialog.getOpenFileName(
        window,
        "Insert image",
        str(Path.home() / "Pictures"),
        "Images (*.png *.jpg *.jpeg *.gif *.bmp *.webp)",
    )
    if not file_name:
        return
    image = QImage(file_name)
    if image.isNull():
        QMessageBox.warning(window, "Insert image", "That picture couldn't be read.")
        return
    if image.width() > 2000:
        image = image.scaledToWidth(2000, Qt.TransformationMode.SmoothTransformation)
    is_photo = Path(file_name).suffix.lower() in (".jpg", ".jpeg")
    data = QByteArray()
    buffer = QBuffer(data)
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    image.save(buffer, "JPG" if is_photo else "PNG", 88 if is_photo else -1)
    buffer.close()
    kind = "image/jpeg" if is_photo else "image/png"
    uri = f"data:{kind};base64,{base64.b64encode(bytes(data)).decode()}"
    editor.document().addResource(image_resource_type(), QUrl(uri), image)

    text_width, _ = page_text_size()
    width = min(float(image.width()), text_width)
    picture = QTextImageFormat()
    picture.setName(uri)
    picture.setWidth(width)
    picture.setHeight(image.height() * width / image.width())
    editor.textCursor().insertImage(picture)
    editor.setFocus()


def image_cursor_at(cursor):
    """A cursor selecting the picture next to this spot, or None."""
    document = editor.document()
    for start in (cursor.position() - 1, cursor.position()):
        if start < 0 or start >= document.characterCount() - 1:
            continue
        picked = QTextCursor(document)
        picked.setPosition(start)
        picked.setPosition(start + 1, QTextCursor.MoveMode.KeepAnchor)
        if picked.charFormat().isImageFormat():
            return picked
    return None


def resize_image(picked, fraction):
    picture = picked.charFormat().toImageFormat()
    ratio = picture.height() / picture.width() if picture.width() else 0.75
    if not ratio:
        image = editor.document().resource(image_resource_type(), QUrl(picture.name()))
        ratio = image.height() / image.width() if image is not None and image.width() else 0.75
    text_width, _ = page_text_size()
    picture.setWidth(text_width * fraction)
    picture.setHeight(text_width * fraction * ratio)
    picked.setCharFormat(picture)


def link_range_at(cursor):
    """(start, end, address) of the link under the cursor, or None."""
    block = cursor.block()
    offset = cursor.positionInBlock()
    for part in block.textFormats():
        if part.format.isAnchor() and part.format.anchorHref() and part.start <= offset <= part.start + part.length:
            return block.position() + part.start, block.position() + part.start + part.length, part.format.anchorHref()
    return None


def normalize_address(address):
    address = address.strip()
    if not address:
        return ""
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", address):
        return address
    if "@" in address and " " not in address and "/" not in address:
        return "mailto:" + address
    return "https://" + address


def insert_link():
    cursor = editor.textCursor()
    existing = link_range_at(cursor) if not cursor.hasSelection() else None
    if existing:
        cursor.setPosition(existing[0])
        cursor.setPosition(existing[1], QTextCursor.MoveMode.KeepAnchor)
        editor.setTextCursor(cursor)

    dialog = QDialog(window)
    dialog.setWindowTitle("Insert link")
    form = QFormLayout(dialog)
    text_box = QLineEdit(cursor.selectedText().replace(" ", " "))
    address_box = QLineEdit(existing[2] if existing else "")
    address_box.setPlaceholderText("Web address or email")
    form.addRow("Text:", text_box)
    form.addRow("Link:", address_box)
    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
    buttons.accepted.connect(dialog.accept)
    buttons.rejected.connect(dialog.reject)
    form.addRow(buttons)
    (address_box if text_box.text() else text_box).setFocus()
    if not dialog.exec():
        return
    address = normalize_address(address_box.text())
    if not address:
        return
    look = QTextCharFormat()
    look.setAnchor(True)
    look.setAnchorHref(address)
    look.setForeground(QColor("#1155cc"))
    look.setFontUnderline(True)
    text = text_box.text() or address
    cursor = editor.textCursor()
    cursor.beginEditBlock()
    if cursor.hasSelection() and cursor.selectedText() == text:
        cursor.mergeCharFormat(look)
    else:
        typed = QTextCharFormat(cursor.charFormat())
        typed.merge(look)
        cursor.insertText(text, typed)
    cursor.endEditBlock()
    editor.setTextCursor(cursor)
    stop_link_typing()


def stop_link_typing():
    """So what you type next isn't part of the link."""
    plain = QTextCharFormat(editor.currentCharFormat())
    if plain.isAnchor():
        plain.setAnchor(False)
        plain.clearProperty(QTextFormat.Property.AnchorHref)
        plain.clearForeground()
        plain.setFontUnderline(False)
        editor.setCurrentCharFormat(plain)


def remove_link(start, end):
    cursor = QTextCursor(editor.document())
    cursor.beginEditBlock()
    position = start
    while position < end:
        piece = QTextCursor(editor.document())
        piece.setPosition(position)
        piece.setPosition(position + 1, QTextCursor.MoveMode.KeepAnchor)
        plain = piece.charFormat()
        plain.setAnchor(False)
        plain.clearProperty(QTextFormat.Property.AnchorHref)
        plain.clearForeground()
        plain.setFontUnderline(False)
        piece.setCharFormat(plain)
        position += 1
    cursor.endEditBlock()


def insert_horizontal_line():
    cursor = editor.textCursor()
    cursor.insertHtml("<hr/>")
    editor.setFocus()


def toggle_full_screen():
    if window.isFullScreen():
        window.showNormal()
    else:
        window.showFullScreen()


def menu_shortcuts(menu, path=""):
    found = []
    for action in menu.actions():
        if action.menu() is not None:
            found += menu_shortcuts(action.menu(), f"{path}{action.text().replace('&&', '&')} → ")
        elif action.text() and not action.isSeparator():
            keys = ", ".join(k.toString(QKeySequence.SequenceFormat.NativeText) for k in action.shortcuts())
            if keys:
                found.append((path + action.text().replace("&&", "&").split("\t")[0], keys))
    return found


def show_shortcuts():
    rows = []
    for top in menu_bar.actions():
        if top.menu() is not None:
            rows += menu_shortcuts(top.menu(), f"{top.text()} → ")
    rows += [
        ("Line break", "Shift+Enter"),
        ("Move a list item in / out", "Tab / Shift+Tab"),
        ("Open a link", "Ctrl+click"),
    ]
    dialog = QDialog(window)
    dialog.setWindowTitle("Keyboard shortcuts")
    column = QVBoxLayout(dialog)
    table = QTableWidget(len(rows), 2)
    table.setHorizontalHeaderLabels(["Action", "Shortcut"])
    table.verticalHeader().setVisible(False)
    table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
    table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
    table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
    for row, (name, keys) in enumerate(rows):
        table.setItem(row, 0, QTableWidgetItem(name))
        table.setItem(row, 1, QTableWidgetItem(keys))
    column.addWidget(table)
    buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
    buttons.rejected.connect(dialog.reject)
    column.addWidget(buttons)
    dialog.resize(560, 600)
    dialog.exec()


def show_about():
    QMessageBox.about(
        window,
        "About Tabbed Writer",
        f"<h3>Tabbed Writer {APP_VERSION}</h3>"
        "<p>A simple, fully offline word processor where each tab is a chapter.</p>"
        "<p>Free software under the GNU General Public License, version 3.<br>"
        "Made with the help of Claude, an AI assistant by Anthropic.</p>"
        '<p><a href="https://github.com/GTRoid76/tabbed-writer">github.com/GTRoid76/tabbed-writer</a></p>',
    )


# ---------------------------------------------------------------------------
# Window
# ---------------------------------------------------------------------------

class Page(QTextEdit):
    """The writing area. Tab and Shift+Tab move list items in and out, like Google Docs."""

    def keyPressEvent(self, event):
        key = enum_number(event.key())
        tab_keys = (enum_number(Qt.Key.Key_Tab), enum_number(Qt.Key.Key_Backtab))
        if key in tab_keys and self.textCursor().currentList():
            backwards = key == enum_number(Qt.Key.Key_Backtab) or bool(
                event.modifiers() & Qt.KeyboardModifier.ShiftModifier
            )
            change_indent(-1 if backwards else 1)
            return
        super().keyPressEvent(event)

    def mouseReleaseEvent(self, event):
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            address = self.anchorAt(event.position().toPoint())
            if address:
                QDesktopServices.openUrl(QUrl(address))
                return
        super().mouseReleaseEvent(event)


class MainWindow(QMainWindow):
    def closeEvent(self, event):
        if ok_to_discard():
            clear_recovery()
            event.accept()
        else:
            event.ignore()


class TabList(QTreeWidget):
    def dropEvent(self, event):
        if self.dropIndicatorPosition() == QAbstractItemView.DropIndicatorPosition.OnItem:
            event.ignore()
            return
        super().dropEvent(event)
        update_counts_soon()


app = QApplication(sys.argv)
app.setDesktopFileName("tabbed-writer")
for icon_file in (
    Path(__file__).with_name("tabbed-writer.svg"),
    Path.home() / ".local/share/icons/hicolor/scalable/apps/tabbed-writer.svg",
):
    if icon_file.exists():
        app.setWindowIcon(QIcon(str(icon_file)))
        break

# Documents written on one system show the matching fonts on the other.
for label in LINUX_FONTS:
    if ON_LINUX:
        QFont.insertSubstitution(OTHER_FONTS[label], LINUX_FONTS[label])
    else:
        QFont.insertSubstitution(LINUX_FONTS[label], OTHER_FONTS[label])

window = MainWindow()
window.setWindowTitle("Untitled - Tabbed Writer")

sections = TabList()
sections.setRootIsDecorated(False)
sections.setHeaderHidden(True)
sections.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)

editor = Page()
editor.setPlaceholderText("Type your document here...")
editor.document().setDefaultFont(QFont(DEFAULT_FONT, DEFAULT_SIZE))
editor.setObjectName("page")
editor.setFrameShape(QFrame.Shape.NoFrame)
editor.setMinimumWidth(420)
editor.setStyleSheet("#page { background: white; color: black; }")
apply_page_look()

page_area = QWidget()
page_area.setObjectName("pageArea")
page_area.setStyleSheet("#pageArea { background: #dadce0; }")
page_row = QHBoxLayout(page_area)
page_row.setContentsMargins(16, 16, 16, 0)
page_row.addStretch()
page_row.addWidget(editor, 100)
page_row.addStretch()
editor.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
editor.customContextMenuRequested.connect(show_editor_menu)
highlighter = SpellHighlighter(editor.document())

plus_button = QPushButton("+")
plus_button.setFixedWidth(30)
plus_button.setToolTip("Add tab (each tab is a chapter)")
plus_button.clicked.connect(add_main_tab)

close_tabs_button = QPushButton("←")
close_tabs_button.setFixedWidth(30)
close_tabs_button.setToolTip("Hide tabs")
close_tabs_button.clicked.connect(lambda: show_tabs_action.setChecked(False))

tabs_header = QHBoxLayout()
tabs_header.addWidget(close_tabs_button)
tabs_header.addWidget(QLabel("Document tabs"))
tabs_header.addStretch()
tabs_header.addWidget(plus_button)

word_panel = QWidget()
word_panel_layout = QVBoxLayout(word_panel)
word_panel_layout.setContentsMargins(0, 6, 0, 0)
word_panel_title = QLabel("Word count")
title_font = word_panel_title.font()
title_font.setBold(True)
word_panel_title.setFont(title_font)
panel_table = make_word_table()
panel_total = QLabel()
panel_total.setFont(title_font)
word_panel_layout.addWidget(word_panel_title)
word_panel_layout.addWidget(panel_table)
word_panel_layout.addWidget(panel_total)
word_panel.setVisible(False)

sidebar_layout = QVBoxLayout()
sidebar_layout.setContentsMargins(0, 0, 0, 0)
sidebar_layout.addLayout(tabs_header)
sidebar_layout.addWidget(sections, 3)
sidebar_layout.addWidget(word_panel, 2)

sidebar = QWidget()
sidebar.setLayout(sidebar_layout)

splitter = QSplitter(Qt.Orientation.Horizontal)
splitter.addWidget(sidebar)
splitter.addWidget(page_area)
splitter.setSizes([220, 780])

layout = QVBoxLayout()
layout.setContentsMargins(0, 0, 0, 0)
layout.addWidget(splitter)

container = QWidget()
container.setLayout(layout)
window.setCentralWidget(container)

count_label = QLabel()
window.statusBar().addPermanentWidget(count_label)
count_timer = QTimer()
count_timer.setSingleShot(True)
count_timer.setInterval(400)
count_timer.timeout.connect(update_counts)

first_tab = QTreeWidgetItem(sections, ["Tab 1"])
first_tab.setData(0, CONTENT_ROLE, "")
first_tab.setData(0, WORDS_ROLE, 0)

sections.currentItemChanged.connect(switch_section)
sections.setCurrentItem(first_tab)

saved_snapshot = current_snapshot()


def show_tab_menu(position):
    item = sections.itemAt(position)
    if item is None:
        return
    sections.setCurrentItem(item)
    tab_menu = QMenu(sections)
    tab_menu.addAction("Rename", rename_section)
    tab_menu.addSeparator()
    tab_menu.addAction("Delete", remove_section)
    tab_menu.exec(sections.viewport().mapToGlobal(position))


sections.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
sections.customContextMenuRequested.connect(show_tab_menu)
editor.textChanged.connect(update_counts_soon)
editor.cursorPositionChanged.connect(update_counts_soon)


# ---------------------------------------------------------------------------
# Text formatting
# ---------------------------------------------------------------------------

def toggle_bold():
    if editor.currentCharFormat().font().bold():
        editor.setFontWeight(QFont.Weight.Normal)
    else:
        editor.setFontWeight(QFont.Weight.Bold)


def toggle_italic():
    editor.setFontItalic(not editor.fontItalic())


def toggle_underline():
    editor.setFontUnderline(not editor.fontUnderline())


def set_alignment(alignment):
    editor.setAlignment(alignment)


def choose_text_color():
    color = choose_color("Text color", editor.textColor())
    if color is not None:
        editor.setTextColor(color)
    editor.setFocus()


def choose_highlight_color():
    color = choose_color("Highlight color", QColor("#fdf3c4"))
    if color is not None:
        highlight = QTextCharFormat()
        highlight.setBackground(color)
        editor.mergeCurrentCharFormat(highlight)
    editor.setFocus()


def remove_colors():
    cursor = editor.textCursor()
    if not cursor.hasSelection():
        plain = editor.currentCharFormat()
        plain.clearForeground()
        plain.clearBackground()
        editor.setCurrentCharFormat(plain)
        return
    start, end = cursor.selectionStart(), cursor.selectionEnd()
    piece = QTextCursor(editor.document())
    piece.beginEditBlock()
    position = start
    while position < end:
        piece.setPosition(position)
        piece.movePosition(QTextCursor.MoveOperation.NextCharacter, QTextCursor.MoveMode.KeepAnchor)
        plain = piece.charFormat()
        plain.clearForeground()
        plain.clearBackground()
        piece.setCharFormat(plain)
        position += 1
    piece.endEditBlock()
    editor.setFocus()


style_box = QComboBox()
style_box.addItems(list(PARAGRAPH_STYLES))
style_box.setToolTip("Paragraph style")
style_box.textActivated.connect(apply_paragraph_style)
editor.cursorPositionChanged.connect(update_style_box)

font_box = QComboBox()
font_box.addItems(list(FONT_CHOICES))
font_box.setToolTip("Font")

size_box = QComboBox()
size_box.addItems(SIZE_CHOICES)
size_box.setCurrentText(str(DEFAULT_SIZE))
size_box.setToolTip("Font size")


def set_font_choice(label):
    editor.setFontFamily(FONT_CHOICES[label])
    editor.setFocus()


def set_size_choice(size):
    editor.setFontPointSize(float(size))
    editor.setFocus()


def update_format_boxes(text_format):
    font_box.setCurrentText(FONT_LABELS.get(text_format.font().family(), "Serif"))
    size = text_format.fontPointSize()
    if size > 0:
        size_box.setCurrentText(str(int(size)))
    else:
        size_box.setCurrentText(str(DEFAULT_SIZE))


font_box.textActivated.connect(set_font_choice)
size_box.textActivated.connect(set_size_choice)
editor.currentCharFormatChanged.connect(update_format_boxes)


# ---------------------------------------------------------------------------
# Menus, with the formatting controls on the same line
# ---------------------------------------------------------------------------

def start_new_window(file_name=None):
    # Starts another copy of Tabbed Writer in its own window.
    if getattr(sys, "frozen", False):
        command = [sys.executable]
    else:
        command = [sys.executable, str(Path(__file__).resolve())]
    if file_name:
        command.append(file_name)
    subprocess.Popen(command)


def open_in_new_window():
    file_name, _ = QFileDialog.getOpenFileName(
        window,
        "Open in new window",
        str(WRITING_FOLDER),
        "All supported (*.twp *.docx *.odt *.txt);;Tabbed Writer document (*.twp);;"
        "Word document (*.docx);;OpenDocument text (*.odt);;Plain text (*.txt)",
    )
    if file_name:
        start_new_window(file_name)


def add_menu_item(menu, name, action_function, shortcut=None):
    action = QAction(name, window)
    if shortcut is not None:
        action.setShortcut(shortcut)
    action.triggered.connect(action_function)
    menu.addAction(action)
    return action


menu_bar = QMenuBar()

file_menu = menu_bar.addMenu("File")
add_menu_item(file_menu, "New", lambda: new_document("standard"), QKeySequence.StandardKey.New)
add_menu_item(file_menu, "New eBook", lambda: new_document("ebook"))
add_menu_item(file_menu, "Open…", open_document, QKeySequence.StandardKey.Open)
recent_menu = file_menu.addMenu("Open recent")
recent_menu.aboutToShow.connect(fill_recent_menu)
add_menu_item(file_menu, "Open in new window…", open_in_new_window)
add_menu_item(file_menu, "New window", start_new_window, QKeySequence("Ctrl+Shift+N"))
add_menu_item(file_menu, "Save", save_document, QKeySequence.StandardKey.Save)
add_menu_item(file_menu, "Save as…", save_as_document, QKeySequence.StandardKey.SaveAs)
add_menu_item(file_menu, "Version history…", show_version_history, QKeySequence("Ctrl+Alt+Shift+H"))
file_menu.addSeparator()
add_menu_item(file_menu, "Page setup…", page_setup)
add_menu_item(file_menu, "Print…", print_document, QKeySequence.StandardKey.Print)
file_menu.addSeparator()
add_menu_item(file_menu, "Close", lambda: new_document("standard"), QKeySequence.StandardKey.Close)
add_menu_item(file_menu, "Exit", window.close)

edit_menu = menu_bar.addMenu("Edit")
add_menu_item(edit_menu, "Undo", editor.undo, QKeySequence.StandardKey.Undo)
add_menu_item(edit_menu, "Redo", editor.redo, QKeySequence.StandardKey.Redo)
edit_menu.addSeparator()
add_menu_item(edit_menu, "Cut", editor.cut, QKeySequence.StandardKey.Cut)
add_menu_item(edit_menu, "Copy", editor.copy, QKeySequence.StandardKey.Copy)
add_menu_item(edit_menu, "Paste", editor.paste, QKeySequence.StandardKey.Paste)
add_menu_item(edit_menu, "Paste without formatting", paste_without_formatting, QKeySequence("Ctrl+Shift+V"))
edit_menu.addSeparator()
add_menu_item(edit_menu, "Select all", editor.selectAll, QKeySequence.StandardKey.SelectAll)
edit_menu.addSeparator()
add_menu_item(edit_menu, "Find…", lambda: show_find(False), QKeySequence("Ctrl+F"))
add_menu_item(edit_menu, "Find and replace…", lambda: show_find(True), QKeySequence("Ctrl+H"))

view_menu = menu_bar.addMenu("View")
show_tabs_action = QAction("Show tabs", window)
show_tabs_action.setCheckable(True)
show_tabs_action.setChecked(True)
show_tabs_action.toggled.connect(sidebar.setVisible)
view_menu.addAction(show_tabs_action)


def set_word_count_panel(turned_on):
    word_panel.setVisible(turned_on)
    app_settings.setValue("show_word_count", turned_on)
    if turned_on:
        show_tabs_action.setChecked(True)
        update_counts()


show_word_count_action = QAction("Show word count", window)
show_word_count_action.setCheckable(True)
show_word_count_action.toggled.connect(set_word_count_panel)
view_menu.addAction(show_word_count_action)
view_menu.addSeparator()
add_menu_item(view_menu, "Full screen", toggle_full_screen, QKeySequence("F11"))

insert_menu = menu_bar.addMenu("Insert")
add_menu_item(insert_menu, "Image…", insert_image)
add_menu_item(insert_menu, "Table…", insert_table)
add_menu_item(insert_menu, "Link…", insert_link, QKeySequence("Ctrl+K"))
add_menu_item(insert_menu, "Horizontal line", insert_horizontal_line)
break_menu = insert_menu.addMenu("Break")
add_menu_item(break_menu, "Page break", insert_page_break, QKeySequence("Ctrl+Return"))
add_menu_item(break_menu, "Line break\tShift+Enter", insert_line_break)

format_menu = menu_bar.addMenu("Format")
text_menu = format_menu.addMenu("Text")
bold_action = add_menu_item(text_menu, "Bold", toggle_bold, QKeySequence.StandardKey.Bold)
italic_action = add_menu_item(text_menu, "Italic", toggle_italic, QKeySequence.StandardKey.Italic)
underline_action = add_menu_item(text_menu, "Underline", toggle_underline, QKeySequence.StandardKey.Underline)
strike_action = add_menu_item(text_menu, "Strikethrough", toggle_strikethrough, QKeySequence("Alt+Shift+5"))
text_menu.addSeparator()
add_menu_item(text_menu, "Superscript", lambda: toggle_vertical(QTextCharFormat.VerticalAlignment.AlignSuperScript), QKeySequence("Ctrl+."))
add_menu_item(text_menu, "Subscript", lambda: toggle_vertical(QTextCharFormat.VerticalAlignment.AlignSubScript), QKeySequence("Ctrl+,"))

styles_menu = format_menu.addMenu("Paragraph styles")
for number, style_name in enumerate(PARAGRAPH_STYLES):
    keys = {"Normal text": "Ctrl+Alt+0", "Heading 1": "Ctrl+Alt+1", "Heading 2": "Ctrl+Alt+2", "Heading 3": "Ctrl+Alt+3"}
    shortcut = QKeySequence(keys[style_name]) if style_name in keys else None
    add_menu_item(styles_menu, style_name, lambda checked=False, n=style_name: apply_paragraph_style(n), shortcut)

align_menu = format_menu.addMenu("Align && indent")
add_menu_item(align_menu, "Left", lambda: set_alignment(Qt.AlignmentFlag.AlignLeft), QKeySequence("Ctrl+Shift+L"))
add_menu_item(align_menu, "Center", lambda: set_alignment(Qt.AlignmentFlag.AlignHCenter), QKeySequence("Ctrl+Shift+E"))
add_menu_item(align_menu, "Right", lambda: set_alignment(Qt.AlignmentFlag.AlignRight), QKeySequence("Ctrl+Shift+R"))
add_menu_item(align_menu, "Justified", lambda: set_alignment(Qt.AlignmentFlag.AlignJustify), QKeySequence("Ctrl+Shift+J"))
align_menu.addSeparator()
indent_more_action = add_menu_item(align_menu, "Increase indent", lambda: change_indent(1), QKeySequence("Ctrl+]"))
indent_less_action = add_menu_item(align_menu, "Decrease indent", lambda: change_indent(-1), QKeySequence("Ctrl+["))

spacing_menu = format_menu.addMenu("Line spacing")
for label, percent in (("Single", 100), ("1.15", 115), ("1.5", 150), ("Double", 200)):
    add_menu_item(spacing_menu, label, lambda checked=False, p=percent: set_line_spacing(p))

lists_menu = format_menu.addMenu("Bullets && numbering")
bullets_action = add_menu_item(lists_menu, "Bulleted list", lambda: toggle_list(QTextListFormat.Style.ListDisc))
bullets_action.setShortcuts([QKeySequence("Ctrl+Shift+8"), QKeySequence("Ctrl+*")])
numbers_action = add_menu_item(lists_menu, "Numbered list", lambda: toggle_list(QTextListFormat.Style.ListDecimal))
numbers_action.setShortcuts([QKeySequence("Ctrl+Shift+7"), QKeySequence("Ctrl+&")])
format_menu.addSeparator()
add_menu_item(format_menu, "Clear formatting", clear_formatting, QKeySequence("Ctrl+\\"))
color_menu = format_menu.addMenu("Color")
text_color_action = add_menu_item(color_menu, "Text color…", choose_text_color)
highlight_action = add_menu_item(color_menu, "Highlight color…", choose_highlight_color)
add_menu_item(color_menu, "Remove colors", remove_colors)
format_menu.addSeparator()
ebook_action = QAction("eBook formatting (KDP / B&N)", window)
ebook_action.setCheckable(True)
ebook_action.toggled.connect(toggle_ebook_formatting)
format_menu.addAction(ebook_action)

tools_menu = menu_bar.addMenu("Tools")
spell_action = QAction("Spell check as you type", window)
spell_action.setCheckable(True)
spell_action.toggled.connect(set_spell_check)
tools_menu.addAction(spell_action)
add_menu_item(tools_menu, "Word count", show_word_count, QKeySequence("Ctrl+Shift+C"))
add_menu_item(tools_menu, "eBook words per page…", change_words_per_page)
tools_menu.addSeparator()
add_menu_item(tools_menu, "Preferences…", show_preferences)

help_menu = menu_bar.addMenu("Help")
add_menu_item(help_menu, "Keyboard shortcuts", show_shortcuts, QKeySequence("Ctrl+/"))
add_menu_item(help_menu, "About Tabbed Writer", show_about)

top_bar = QWidget()
top_row = QHBoxLayout(top_bar)
top_row.setContentsMargins(0, 0, 6, 0)
top_row.setSpacing(4)
top_row.addWidget(menu_bar)
top_row.addSpacing(12)
top_row.addWidget(style_box)
top_row.addWidget(font_box)
top_row.addWidget(size_box)

for action, letter, style in [
    (bold_action, "B", "bold"),
    (italic_action, "I", "italic"),
    (underline_action, "U", "underline"),
]:
    action.setIconText(letter)
    action.setToolTip(action.text())
    button = QToolButton()
    button.setDefaultAction(action)
    button.setAutoRaise(True)
    button_font = QFont()
    if style == "bold":
        button_font.setBold(True)
    elif style == "italic":
        button_font.setItalic(True)
    else:
        button_font.setUnderline(True)
    button.setFont(button_font)
    top_row.addWidget(button)

for action, letter, look in [
    (text_color_action, "A", "QToolButton { color: #c62828; font-weight: bold; }"),
    (highlight_action, "A", "QToolButton { background: #fdf3c4; color: black; font-weight: bold; }"),
]:
    action.setIconText(letter)
    action.setToolTip(action.text().replace("…", ""))
    button = QToolButton()
    button.setDefaultAction(action)
    button.setAutoRaise(True)
    button.setStyleSheet(look)
    top_row.addWidget(button)

for action, label in [
    (bullets_action, "•"),
    (numbers_action, "1."),
    (indent_less_action, "⇤"),
    (indent_more_action, "⇥"),
]:
    action.setIconText(label)
    action.setToolTip(action.text())
    button = QToolButton()
    button.setDefaultAction(action)
    button.setAutoRaise(True)
    top_row.addWidget(button)

top_row.addStretch()
window.setMenuWidget(top_bar)

show_word_count_action.setChecked(
    str(app_settings.value("show_word_count", "false")).lower() == "true"
)
update_counts()
if len(sys.argv) > 1:
    open_document(sys.argv[1])
window.resize(1100, 700)
window.show()
editor.setFocus()

spell_on = str(app_settings.value("spell_check", "true")).lower() == "true"
spell_action.blockSignals(True)
spell_action.setChecked(spell_on)
spell_action.blockSignals(False)
highlighter.enabled = spell_on
QTimer.singleShot(150, load_spell_checker)

autosave_timer = QTimer()
autosave_timer.setInterval(60 * 1000)
autosave_timer.timeout.connect(autosave_tick)
autosave_timer.start()
QTimer.singleShot(600, check_for_recovery)

sys.exit(app.exec())
